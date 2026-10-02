"""A report as a file: CSV, a workbook, or a PDF.

Three writers over one result shape, chosen by `ExportFormat`. CSV goes
through the standard library's writer with one defence a writer does not
give: a cell that begins with a character a spreadsheet reads as a formula
is prefixed with an apostrophe, so a tenant's own data -- a file named
`=HYPERLINK(...)`, a member called `+1` -- opens as text and executes
nothing on the machine of whoever downloaded it. The workbook writer
stores such a cell as text for the same reason. The PDF is a page of
figures and a table of rows, in the framework's own formatting, and no
branding: a product that wants its mark on a document renders its own.

Bounded. A result with more rows than `EXPORT_ROW_LIMIT` is handed to the
background export by the caller before it gets here, or -- where the caller
has no background to hand it to, which is a scheduled delivery -- refused;
these write what they are given.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from .definitions import ExportFormat, Format
from .results import Cell, MetricValue, ReportResult

#: The most rows one synchronous export carries. Past this the export runs
#: in the background and lands in the tenant's bucket.
#:
#: A handoff threshold, and not the most an export may be: the background
#: export has no row maximum of its own.
#:
#: The worker's scheduled delivery derives its maximum from this number, as
#: `SCHEDULED_DELIVERY_ROW_LIMIT` in `koras_worker/tasks/reporting.py`, since
#: GR-352E. There it *is* a maximum: a delivery is an attachment built whole
#: in the worker and has no bucket to fall back to, so past it the delivery is
#: refused and the schedule says why. One number, on purpose, and two
#: meanings, each under its own name.
EXPORT_ROW_LIMIT = 10_000

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


@dataclass(frozen=True)
class Rendered:
    content: bytes
    media_type: str
    extension: str


MEDIA_TYPES: dict[ExportFormat, tuple[str, str]] = {
    ExportFormat.CSV: ("text/csv; charset=utf-8", "csv"),
    ExportFormat.XLSX: (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
    ExportFormat.PDF: ("application/pdf", "pdf"),
}


def _text(value: Cell) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # A count that travelled as a float is still a count on paper.
        value = int(value)
    return str(value)


def _cell(value: Cell) -> str:
    text = _text(value)
    if text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def _columns_and_rows(result: ReportResult) -> tuple[list[str], list[list[Cell]]]:
    """The table when there is one, else the metrics as rows."""
    if result.table is not None and result.table.columns:
        labels = [column.label for column in result.table.columns]
        rows = [
            [row.get(column.key) for column in result.table.columns] for row in result.table.rows
        ]
        return labels, rows
    return ["metric", "value", "unit", "kind"], [
        [metric.label, metric.value, metric.unit.value, metric.kind] for metric in result.metrics
    ]


def to_csv(result: ReportResult) -> str:
    """The rows of the report's table, or its metrics when it has no table."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    labels, rows = _columns_and_rows(result)
    writer.writerow([_cell(label) for label in labels])
    for row in rows:
        writer.writerow([_cell(value) for value in row])
    return buffer.getvalue()


def to_xlsx(result: ReportResult) -> bytes:
    """One sheet of rows, and a second of the figures, as a workbook."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active if book.active is not None else book.create_sheet("Rows")
    sheet.title = "Rows"
    labels, rows = _columns_and_rows(result)
    sheet.append(labels)
    for row in rows:
        sheet.append([_workbook_cell(value) for value in row])
    if result.table is not None and result.metrics:
        figures = book.create_sheet("Figures")
        figures.append(["metric", "value", "unit", "kind", "limit", "previous"])
        for metric in result.metrics:
            figures.append(
                [
                    metric.label,
                    _workbook_cell(metric.value),
                    metric.unit.value,
                    metric.kind,
                    _workbook_cell(metric.limit),
                    _workbook_cell(metric.previous),
                ]
            )
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _workbook_cell(value: Cell) -> Cell:
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        # A string cell, never a formula: openpyxl would otherwise store an
        # `=` cell as one.
        return "'" + value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def to_pdf(result: ReportResult, *, title: str | None = None) -> bytes:
    """A page of figures and a table of rows. Landscape, so a wide report fits."""
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(0, 9, _latin(title or result.key), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    if result.range is not None:
        pdf.cell(0, 6, f"{result.range.start} to {result.range.end}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Generated {result.generated_at.isoformat()}", new_x="LMARGIN", new_y="NEXT")
    for note in result.notes:
        pdf.multi_cell(0, 5, _latin(note))
    pdf.ln(2)

    if result.metrics:
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 7, "Figures", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=9)
        for metric in result.metrics:
            pdf.cell(0, 5, _latin(_figure_line(metric)), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    labels, rows = _columns_and_rows(result)
    if result.table is not None and rows:
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 7, "Rows", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=8)
        usable = pdf.w - pdf.l_margin - pdf.r_margin
        width = usable / max(1, len(labels))
        pdf.set_font("Helvetica", "B", 8)
        for label in labels:
            pdf.cell(width, 6, _latin(_clip(label)), border=1)
        pdf.ln(6)
        pdf.set_font("Helvetica", size=8)
        for row in rows:
            for value in row:
                pdf.cell(width, 5, _latin(_clip(_text(value))), border=1)
            pdf.ln(5)
    return bytes(pdf.output())


def _figure_line(metric: MetricValue) -> str:
    value = _text(metric.value) if metric.value is not None else "not available"
    parts = [f"{metric.label}: {value}"]
    if metric.format is Format.MONEY and metric.value is not None:
        parts = [f"{metric.label}: ${metric.value / 1_000_000:,.2f}"]
    if metric.limit is not None:
        parts.append(f"of {_text(metric.limit)}")
    if metric.kind != "actual":
        parts.append(f"({metric.kind})")
    if metric.note:
        parts.append(f"- {metric.note}")
    return " ".join(parts)


def _clip(text: str, width: int = 28) -> str:
    return text if len(text) <= width else text[: width - 1] + "…"


def _latin(text: str) -> str:
    """The core PDF fonts speak Latin-1; anything else becomes a question mark
    rather than a crash, and a product with other scripts embeds a font."""
    return text.encode("latin-1", "replace").decode("latin-1")


def render(result: ReportResult, fmt: ExportFormat, *, title: str | None = None) -> Rendered:
    media_type, extension = MEDIA_TYPES[fmt]
    if fmt is ExportFormat.CSV:
        return Rendered(to_csv(result).encode("utf-8"), media_type, extension)
    if fmt is ExportFormat.XLSX:
        return Rendered(to_xlsx(result), media_type, extension)
    return Rendered(to_pdf(result, title=title), media_type, extension)


def export_filename(key: str, start: str | None, end: str | None, extension: str = "csv") -> str:
    stem = key.replace(".", "-")
    if start and end:
        return f"{stem}-{start}-to-{end}.{extension}"
    return f"{stem}.{extension}"


def row_count(result: ReportResult) -> int:
    return len(result.table.rows) if result.table is not None else len(result.metrics)
