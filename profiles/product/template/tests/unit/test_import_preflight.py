"""No path to a reader that does not go through the safety pass first.

GR-352. `koras-import` has its own suite for what the pass counts and refuses.
What only exists once the engine is wired into this product is the *order*:
four functions in `core/imports.py` take a customer's bytes -- `analyse` for
the two routes, `check`, `examine` and `prepare` for the dry run and the commit
-- and each must refuse an unsafe file before the expensive thing it would
otherwise do with it.

That is asserted by replacing the expensive things with something that fails
the test when reached: `openpyxl.load_workbook` for a workbook and the
whole-file `decode` for a CSV. A second test then shows the same replacements
*are* reached by a file that passes, because a guard that cannot fire proves
nothing about the code it stands in front of.

**Since GR-352B `analyse` has a different expensive thing.** It no longer
calls either reader, for any file: it streams the head through
`koras_import.inspection`. So what stands behind the safety pass there is the
inspection, and that is what is replaced for it -- while the two readers stay
replaced too, which now asserts they are never reached from `analyse` at all.
`test_import_inspection.py` has the rest of that route's assertions.

The limits are made small here so no test builds a large file. The envelope a
product actually runs under is `SAFETY_LIMITS`, and the numbers in it are
provisional.
"""

from __future__ import annotations

import io
import os
import zipfile
from collections.abc import Callable
from dataclasses import replace
from datetime import UTC, datetime
from typing import Any

import pytest

# What importing the routes needs, as `test_audit_actions.py` supplies it. This
# file runs alone as well as in the suite.
os.environ.setdefault("ENVIRONMENT", "dev")
os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://user:pass@localhost/db")
os.environ.setdefault("ZITADEL_DOMAIN", "https://example.invalid")
os.environ.setdefault("ZITADEL_PROJECT_ID", "0")

from koras_api.core import imports as store  # noqa: E402
from koras_api.core.errors import ApiErrorCode  # noqa: E402
from koras_import import (  # noqa: E402
    FieldSpec,
    Format,
    ImportTarget,
    Operation,
    PreflightRefused,
    ReadRefused,
    RefusalCode,
    RunState,
    SafetyLimits,
    inspection,
)

EMOJI = "\U0001f600"

PEOPLE = ImportTarget(
    key="probe.people",
    label_key="import.target.probe.people",
    permission="imports.manage",
    fields=(FieldSpec(name="name", label_key="import.field.name", required=True),),
    match_keys=("name",),
    operations=(Operation.SKIP_DUPLICATE,),
    formats=(Format.CSV, Format.XLSX),
    max_rows=1000,
)

CSV_RUN = store.Run(
    id="11111111-1111-1111-1111-111111111111",
    target=PEOPLE.key,
    status=RunState.MAPPED.value,
    format=Format.CSV.value,
    source_file_id="22222222-2222-2222-2222-222222222222",
    delimiter=",",
    encoding="utf-8",
    columns=("Name",),
    mapping={"Name": "name"},
    operation=Operation.SKIP_DUPLICATE.value,
    rows_total=0,
    rows_valid=0,
    errors_total=0,
    errors_cut=False,
    error=None,
    requested_by="user-1",
    committed_by=None,
    created_at=datetime.now(UTC),
    started_at=None,
    finished_at=None,
)
XLSX_RUN = replace(CSV_RUN, format=Format.XLSX.value)


def workbook(rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    made = Workbook()
    sheet = made.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    made.save(buffer)
    return buffer.getvalue()


#: A file that is ASCII but for one character, which is the GR-352 shape: the
#: reader would hold all of it at four bytes a character.
WIDE_CSV = ("Name\n" + "Ada Lovelace\n" * 2000 + EMOJI + "\n").encode()
WIDE_XLSX = workbook([["Name"], ["a" * 20_000 + EMOJI]])
SAFE_CSV = b"Name\nAda\n"
SAFE_XLSX = workbook([["Name"], ["Ada"]])
TWO_WIDE_XLSX = workbook([["Name", "Email"], ["Ada", "ada@example.com"]])



def _package(parts: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    return buffer.getvalue()


_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
_SHEET_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
_STRINGS_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"


def _odd_workbook(*, strings_root: str = "sst", workbook_part: str = "xl/workbook.xml") -> bytes:
    """A workbook whose text is in a part the safety pass used to walk past.

    IMPORT-DEF-017. Sixty thousand characters, held either in a string table
    whose root element is not `sst`, or in a sheet listed only by a second
    workbook part that the content types name and `xl/workbook.xml` does not.
    `openpyxl` reads both. `test_preflight_package.py` in the engine has every
    such shape; these two are here so the *worker's* three ways in are asked.
    """
    long = "a" * 60_000
    sheet = (
        f'<worksheet xmlns="{_MAIN}"><sheetData>'
        '<row r="1"><c r="A1" t="s"><v>0</v></c></row>'
        '<row r="2"><c r="A2" t="s"><v>1</v></c></row></sheetData></worksheet>'
    )
    small = (
        f'<worksheet xmlns="{_MAIN}"><sheetData><row r="1"><c r="A1" t="inlineStr">'
        "<is><t>Name</t></is></c></row></sheetData></worksheet>"
    )
    listing = (
        f'<workbook xmlns="{_MAIN}" xmlns:r="{_RELS}"><sheets>'
        '<sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )

    def relations(target: str) -> str:
        return (
            f'<Relationships xmlns="{_PKG}"><Relationship Id="rId1" '
            f'Type="{_RELS}/worksheet" Target="{target}"/></Relationships>'
        )

    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/{workbook_part}" ContentType="{_SHEET_TYPE}"/>'
            f'<Override PartName="/xl/sharedStrings.xml" ContentType="{_STRINGS_TYPE}"/></Types>'
        ),
        "xl/workbook.xml": listing,
        "xl/_rels/workbook.xml.rels": relations("worksheets/small.xml"),
        "xl/worksheets/small.xml": small,
        "xl/worksheets/heavy.xml": sheet,
        "xl/sharedStrings.xml": (
            f'<{strings_root} xmlns="{_MAIN}"><si><t>Name</t></si>'
            f"<si><t>{long}</t></si></{strings_root}>"
        ),
    }
    if workbook_part != "xl/workbook.xml":
        parts[workbook_part] = listing
        folder, _, name = workbook_part.rpartition("/")
        parts[f"{folder}/_rels/{name}.rels"] = relations("worksheets/heavy.xml")
    else:
        parts["xl/_rels/workbook.xml.rels"] = relations("worksheets/heavy.xml")
    return _package(parts)


RENAMED_ROOT_XLSX = _odd_workbook(strings_root="strings")
SECOND_WORKBOOK_XLSX = _odd_workbook(workbook_part="xl/real.xml")

#: Between the two: the safe files fit, the wide ones do not.
TIGHT = SafetyLimits(max_decoded_string_bytes=50_000)

#: Every way this module turns a customer's bytes into rows. The worker reaches
#: `examine` and `prepare`; the routes reach `analyse`; `check` is the dry run
#: without a matcher.
READERS: dict[str, Callable[[bytes, Format], Any]] = {
    "analyse": lambda raw, fmt: store.analyse(raw, PEOPLE, fmt),
    "check": lambda raw, fmt: store.check(raw, PEOPLE, _run(fmt)),
    "examine": lambda raw, fmt: store.examine(raw, PEOPLE, _run(fmt)),
    "prepare": lambda raw, fmt: store.prepare(raw, PEOPLE, _run(fmt)),
}

#: The three that parse the whole file through the canonical readers. The
#: fourth, `analyse`, has inspected the head instead since GR-352B.
PARSERS = ("check", "examine", "prepare")


def _run(fmt: Format) -> store.Run:
    return XLSX_RUN if fmt is Format.XLSX else CSV_RUN


@pytest.fixture
def guarded(monkeypatch: pytest.MonkeyPatch) -> None:
    """A tight envelope, and every expensive read replaced with a failure."""
    import openpyxl
    from koras_import import reading

    def workbook_reached(*_: object, **__: object) -> None:
        raise AssertionError("openpyxl was reached before the safety pass refused")

    def decode_reached(_: bytes) -> None:
        raise AssertionError("the whole file was decoded before the safety pass refused")

    def inspection_reached(*_: object, **__: object) -> None:
        raise AssertionError("the file was inspected before the safety pass refused")

    monkeypatch.setattr(store, "SAFETY_LIMITS", TIGHT)
    monkeypatch.setattr(openpyxl, "load_workbook", workbook_reached)
    monkeypatch.setattr(store, "decode", decode_reached)
    monkeypatch.setattr(reading, "decode", decode_reached)
    monkeypatch.setattr(inspection, "_inspect_csv", inspection_reached)
    monkeypatch.setattr(inspection, "_inspect_xlsx", inspection_reached)


@pytest.mark.parametrize("reader", sorted(READERS))
@pytest.mark.parametrize(
    ("raw", "fmt"),
    [
        (WIDE_CSV, Format.CSV),
        (WIDE_XLSX, Format.XLSX),
        # IMPORT-DEF-017: both of these were called safe, and read in full.
        (RENAMED_ROOT_XLSX, Format.XLSX),
        (SECOND_WORKBOOK_XLSX, Format.XLSX),
    ],
    ids=["csv", "xlsx", "xlsx-renamed-root", "xlsx-second-workbook-part"],
)
@pytest.mark.usefixtures("guarded")
def test_an_unsafe_file_is_refused_before_any_reader_is_handed_it(
    reader: str, raw: bytes, fmt: Format
) -> None:
    with pytest.raises(PreflightRefused) as caught:
        READERS[reader](raw, fmt)
    assert caught.value.code is RefusalCode.DECODED_TOO_LARGE


@pytest.mark.parametrize("reader", PARSERS)
@pytest.mark.usefixtures("guarded")
def test_the_guard_fires_for_a_csv_that_passes(reader: str) -> None:
    # Not `ReadRefused`: nothing catches the failure, so it arrives as itself.
    with pytest.raises(AssertionError, match="whole file was decoded"):
        READERS[reader](SAFE_CSV, Format.CSV)


@pytest.mark.parametrize("reader", PARSERS)
@pytest.mark.usefixtures("guarded")
def test_the_guard_fires_for_a_workbook_that_passes(reader: str) -> None:
    # `read_workbook` turns whatever `openpyxl` raises into its own refusal,
    # so the failure is found as the cause.
    with pytest.raises(ReadRefused) as caught:
        READERS[reader](SAFE_XLSX, Format.XLSX)
    assert not isinstance(caught.value, PreflightRefused)
    assert isinstance(caught.value.__cause__, AssertionError)


@pytest.mark.parametrize(
    ("raw", "fmt"), [(SAFE_CSV, Format.CSV), (SAFE_XLSX, Format.XLSX)], ids=["csv", "xlsx"]
)
@pytest.mark.usefixtures("guarded")
def test_the_guard_fires_for_an_analysis_of_a_file_that_passes(raw: bytes, fmt: Format) -> None:
    # The analysis reaches the inspection and neither reader: with all three
    # replaced, the inspection's is the failure that arrives.
    with pytest.raises(AssertionError, match="the file was inspected"):
        READERS["analyse"](raw, fmt)


def test_the_routes_and_the_worker_are_given_one_envelope() -> None:
    for_routes = store.limits_for(PEOPLE, rows=False)
    for_worker = store.limits_for(PEOPLE, rows=True)
    assert replace(for_worker, max_rows=None) == for_routes == store.SAFETY_LIMITS
    assert for_worker.max_rows == PEOPLE.max_rows
    assert store.SAFETY_LIMITS.max_source_bytes == store.MAX_SOURCE_BYTES


@pytest.mark.parametrize("fmt", [Format.CSV, Format.XLSX])
def test_too_many_rows_is_still_the_analysis_answering_and_the_dry_run_refusing(
    fmt: Format,
) -> None:
    small = replace(PEOPLE, max_rows=2)
    rows: list[list[object]] = [["Name"]] + [[f"P{number}"] for number in range(5)]
    raw = (
        workbook(rows)
        if fmt is Format.XLSX
        else "".join(f"{row[0]}\n" for row in rows).encode()
    )
    # The route's own answer is unchanged: a 200 saying the file is too long,
    # which the mapping route then refuses.
    assert store.analyse(raw, small, fmt).over_ceiling is True
    # The dry run and the commit used to read `max_rows` rows of it and say
    # nothing about the rest.
    with pytest.raises(PreflightRefused) as caught:
        store.check(raw, small, _run(fmt))
    assert caught.value.code is RefusalCode.TOO_MANY_ROWS


@pytest.mark.parametrize("fmt", [Format.CSV, Format.XLSX])
def test_a_file_exactly_at_the_row_ceiling_is_read(fmt: Format) -> None:
    exact = replace(PEOPLE, max_rows=3)
    rows: list[list[object]] = [["Name"], ["A"], ["B"], ["C"]]
    raw = (
        workbook(rows)
        if fmt is Format.XLSX
        else b"Name\nA\n\nB\nC\n"  # a blank line is not a row
    )
    assert store.check(raw, exact, _run(fmt)).rows == 3


def test_a_refusal_is_answered_with_a_code_this_api_already_had() -> None:
    from koras_api.routers.imports import _read_refusal

    def answer(code: RefusalCode) -> tuple[int, str]:
        raised = _read_refusal(PreflightRefused(code, "a sentence for a person"))
        return raised.status_code, raised.detail["code"]

    too_large = (413, ApiErrorCode.IMPORT_FILE_TOO_LARGE.value)
    assert answer(RefusalCode.SOURCE_TOO_LARGE) == too_large
    assert answer(RefusalCode.UNCOMPRESSED_TOO_LARGE) == too_large
    assert answer(RefusalCode.DECODED_TOO_LARGE) == too_large
    assert answer(RefusalCode.TOO_MANY_CELLS) == too_large
    assert answer(RefusalCode.TOO_MANY_COLUMNS) == too_large
    assert answer(RefusalCode.TOO_MANY_ENTRIES) == too_large
    assert answer(RefusalCode.TOO_MANY_ROWS) == (422, ApiErrorCode.IMPORT_TOO_MANY_ROWS.value)
    unreadable = (422, ApiErrorCode.IMPORT_FILE_UNREADABLE.value)
    assert answer(RefusalCode.MALFORMED) == unreadable
    assert answer(RefusalCode.MACROS) == unreadable
    assert answer(RefusalCode.LINE_TOO_LONG) == unreadable
    # And a refusal that is the reader's own is what it always was.
    plain = _read_refusal(ReadRefused("the file is empty"))
    assert (plain.status_code, plain.detail["code"]) == unreadable


@pytest.mark.parametrize(
    ("raw", "fmt", "limits"),
    [
        (WIDE_CSV, Format.CSV, TIGHT),
        (WIDE_XLSX, Format.XLSX, TIGHT),
        (SAFE_XLSX, Format.XLSX, SafetyLimits(max_cells=1)),
        (TWO_WIDE_XLSX, Format.XLSX, SafetyLimits(max_columns=1)),
        (b"\x00not a zip", Format.XLSX, SafetyLimits()),
    ],
    ids=["csv-text", "xlsx-text", "cells", "columns", "not-a-workbook"],
)
def test_the_sentence_a_failed_run_records_names_no_internals(
    monkeypatch: pytest.MonkeyPatch, raw: bytes, fmt: Format, limits: SafetyLimits
) -> None:
    # The worker stores `str(refused)` on the run and a person reads it.
    monkeypatch.setattr(store, "SAFETY_LIMITS", limits)
    with pytest.raises(ReadRefused) as caught:
        store.check(raw, PEOPLE, _run(fmt))
    sentence = str(caught.value).lower()
    for internal in ("openpyxl", "expat", "zipfile", "traceback", "memoryerror", "unicode"):
        assert internal not in sentence
    assert 0 < len(sentence) <= 400
