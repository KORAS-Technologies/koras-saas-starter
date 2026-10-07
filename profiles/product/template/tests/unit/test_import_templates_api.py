"""The template route, asked of the app with its dependencies replaced.

Three things a test that reads the router's text cannot prove: that the
permission check happens *before* anything is rendered, that a format the
target does not accept is a 406 and not a file, and that the download leaves
an audit row under the caller. Each is asked of the route here, and the
first is mutation-checked in the generator suite by removing a helper and
expecting red.
"""

from __future__ import annotations

import io
import os
from collections.abc import AsyncIterator, Iterator
from typing import Any

import pytest

os.environ.setdefault("KORAS_DATABASE_URL", "postgresql://x:y@localhost/z")

from fastapi.testclient import TestClient  # noqa: E402
from koras_api.core.auth import require_auth  # noqa: E402
from koras_api.core.database import get_db  # noqa: E402
from koras_api.core.settings import settings as api_settings  # noqa: E402
from koras_api.core.tenant import require_tenant  # noqa: E402
from koras_api.imports import registry  # noqa: E402
from koras_api.main import app  # noqa: E402
from koras_auth import JWTClaims  # noqa: E402
from koras_import import FieldKind, FieldSpec, Format, ImportTarget, Operation  # noqa: E402
from koras_platform import OrganizationRole  # noqa: E402
from koras_tenant import TenantContext  # noqa: E402

app.state.redis = None

TENANT = "00000000-0000-0000-0000-000000000001"
SUBJECT = "user-1"

ACCOUNTS = ImportTarget(
    key="probe.accounts",
    label_key="import.target.probe.accounts",
    permission="imports.manage",
    fields=(
        FieldSpec("email", "import.field.email", kind=FieldKind.EMAIL, required=True),
        FieldSpec("name", "import.field.name", required=True, example="Example Ltd"),
        FieldSpec("tier", "import.field.tier", options=("basic", "pro")),
    ),
    match_keys=("email",),
    operations=(Operation.SKIP_DUPLICATE,),
    formats=(Format.CSV, Format.XLSX),
    version=2,
)

#: A target whose own permission nobody holds, so the target's gate is what
#: refuses an owner who holds `imports.manage` and everything else.
GUARDED = ImportTarget(
    key="probe.secrets",
    label_key="import.target.probe.secrets",
    permission="nobody.holds",
    fields=(FieldSpec("code", "import.field.code", required=True),),
    operations=(Operation.CREATE,),
)


class _Result:
    def first(self) -> None:
        return None

    def one(self) -> None:
        raise AssertionError("nothing here selects a row")

    def __iter__(self) -> Iterator[Any]:
        return iter(())


class _Session:
    def __init__(self) -> None:
        self.statements: list[tuple[str, dict[str, Any]]] = []
        self.commits = 0

    async def execute(self, statement: object, parameters: dict[str, Any] | None = None) -> _Result:
        self.statements.append((" ".join(str(statement).split()), dict(parameters or {})))
        return _Result()

    async def commit(self) -> None:
        self.commits += 1

    async def rollback(self) -> None:
        return None

    def audited(self) -> list[dict[str, Any]]:
        return [params for sql, params in self.statements if "public.audit_events" in sql]


def _install(session: _Session, roles: frozenset[OrganizationRole]) -> TestClient:
    async def _db() -> AsyncIterator[_Session]:
        yield session

    app.dependency_overrides[require_auth] = lambda: JWTClaims(
        sub=SUBJECT, roles=roles, organization_id="org-1"
    )
    app.dependency_overrides[require_tenant] = lambda: TenantContext(
        id=TENANT, slug="alpha", name="Alpha", organization_id="org-1", user_id=SUBJECT
    )
    app.dependency_overrides[get_db] = _db
    return TestClient(app)


@pytest.fixture(autouse=True)
def _targets(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    # Activation gate open: these tests are about templates, not about the switch.
    monkeypatch.setattr(api_settings, "imports_enabled", True)
    registry.clear()
    registry.extend([ACCOUNTS, GUARDED])
    yield
    registry.clear()
    app.dependency_overrides.clear()


OWNER = frozenset({OrganizationRole.OWNER})
MEMBER = frozenset({OrganizationRole.MEMBER})


def test_a_csv_template_is_the_header_row_and_an_audit_row() -> None:
    session = _Session()
    client = _install(session, OWNER)
    answer = client.get("/api/v1/imports/targets/probe.accounts/template?format=csv")
    assert answer.status_code == 200
    assert answer.headers["content-type"].startswith("text/csv")
    assert 'filename="probe-accounts-template-v2.csv"' in answer.headers["content-disposition"]
    assert answer.content.decode("utf-8-sig") == "email,name,tier\r\n"
    rows = session.audited()
    assert len(rows) == 1
    assert rows[0]["action"] == "import.template.downloaded"
    assert rows[0]["actor_id"] == SUBJECT
    assert session.commits == 1


def test_an_xlsx_template_opens_as_a_workbook_with_the_declared_columns() -> None:
    from openpyxl import load_workbook

    client = _install(_Session(), OWNER)
    answer = client.get("/api/v1/imports/targets/probe.accounts/template?format=xlsx")
    assert answer.status_code == 200
    assert "spreadsheetml" in answer.headers["content-type"]
    book = load_workbook(io.BytesIO(answer.content))
    assert [cell.value for cell in book["Data"][1]] == ["email", "name", "tier"]
    assert "Instructions" in book.sheetnames


def test_the_permission_is_checked_before_anything_is_rendered() -> None:
    session = _Session()
    client = _install(session, MEMBER)
    answer = client.get("/api/v1/imports/targets/probe.accounts/template?format=csv")
    assert answer.status_code == 403
    assert session.audited() == []
    assert session.statements == []


def test_the_targets_own_permission_gates_its_template_too() -> None:
    session = _Session()
    client = _install(session, OWNER)
    answer = client.get("/api/v1/imports/targets/probe.secrets/template?format=csv")
    assert answer.status_code == 403
    assert session.audited() == []
    # And the list never offered it, so the page never drew the control.
    listed = client.get("/api/v1/imports/targets").json()
    assert [t["key"] for t in listed] == ["probe.accounts"]


def test_a_format_the_target_does_not_accept_is_406_not_a_file() -> None:
    session = _Session()
    client = _install(session, OWNER)
    for wanted in ("json", "pdf", "docx"):
        answer = client.get(f"/api/v1/imports/targets/probe.accounts/template?format={wanted}")
        assert answer.status_code == 406, wanted
        assert answer.json()["detail"]["code"] == "import_format_refused"
    assert session.audited() == []


def test_an_unknown_target_is_404() -> None:
    client = _install(_Session(), OWNER)
    answer = client.get("/api/v1/imports/targets/probe.nothing/template?format=csv")
    assert answer.status_code == 404


def test_the_target_list_carries_what_the_template_menu_and_the_page_need() -> None:
    client = _install(_Session(), OWNER)
    answer = client.get("/api/v1/imports/targets")
    assert answer.status_code == 200
    [accounts] = [t for t in answer.json() if t["key"] == "probe.accounts"]
    assert accounts["version"] == 2
    assert len(accounts["fingerprint"]) == 12
    assert accounts["template_formats"] == ["csv", "xlsx"]
    assert accounts["limits"]["max_rows"] == accounts["max_rows"]
    assert accounts["limits"]["max_bytes"] <= 64 * 1024 * 1024
    assert accounts["predicts"] is False
    by_name = {field["name"]: field for field in accounts["fields"]}
    assert by_name["name"]["example"] == "Example Ltd"
    assert by_name["email"]["example"] == "name@example.com"
    assert by_name["tier"]["format_hint"] == "one of: basic, pro"
