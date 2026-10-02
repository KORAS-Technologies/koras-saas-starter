"""The analysis and mapping routes, since they stopped parsing the whole file.

GR-352B. Until it, both routes handed a customer's file to the readers the dry
run uses -- `read_workbook`, and `decode` into one string -- inline, on the
event loop, to keep two hundred rows of it. `koras-import` has its own suite
for the inspection that replaced that. What only exists once it is wired into
this product is asked here, and each is something a test reading the router's
text could not prove.

**The answer is the one the routes gave before.** `analysed_as_before` below is
the body `core/imports.analyse` had until GR-352B, kept verbatim as the
reference, and every file is handed to both: the columns, the suggested
mapping, the template verdict, the preview, the count. Two representative
targets, because a mapping suggestion is a function of the target as well as
the header.

**The mapping route reads no rows at all.** It asks for the header and the
count, and reaches neither reader -- proved by replacing both with something
that fails the test.

**Neither route holds the event loop.** The inspection is replaced with one
that blocks until the test releases it, and the test releases it only after
another request has been answered by the same application. Run inline, the
other request cannot be answered, the block times out, and the test fails.
Nothing in it depends on how fast anything is.

**The response is the same shape.** No field was added to say the route
changed, and none was removed.
"""

from __future__ import annotations

import asyncio
import io
import os
import threading
from collections.abc import AsyncIterator, Iterator
from dataclasses import replace
from datetime import UTC, date, datetime
from types import SimpleNamespace
from typing import Any

import pytest

os.environ.setdefault("ENVIRONMENT", "dev")
os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://user:pass@localhost/db")
os.environ.setdefault("KORAS_DATABASE_URL", "postgresql://x:y@localhost/z")
os.environ.setdefault("ZITADEL_DOMAIN", "https://example.invalid")
os.environ.setdefault("ZITADEL_PROJECT_ID", "0")

from fastapi.testclient import TestClient  # noqa: E402
from httpx import ASGITransport, AsyncClient  # noqa: E402
from koras_api.core import imports as store  # noqa: E402
from koras_api.core.auth import require_auth  # noqa: E402
from koras_api.core.database import get_db  # noqa: E402
from koras_api.core.storage import tenant_storage  # noqa: E402
from koras_api.core.tenant import require_tenant  # noqa: E402
from koras_api.imports import registry  # noqa: E402
from koras_api.main import app  # noqa: E402
from koras_auth import JWTClaims  # noqa: E402
from koras_import import (  # noqa: E402
    FieldKind,
    FieldSpec,
    Format,
    ImportTarget,
    Operation,
    ReadRefused,
    RunState,
    SafetyLimits,
    compare,
    count_rows,
    decode,
    preflight,
    read_header,
    read_rows,
    read_workbook,
    render_xlsx,
    sniff_delimiter,
    suggest,
)
from koras_platform import OrganizationRole  # noqa: E402
from koras_tenant import TenantContext  # noqa: E402

app.state.redis = None

TENANT = "00000000-0000-0000-0000-000000000001"
SUBJECT = "user-1"
RUN_ID = "11111111-1111-1111-1111-111111111111"
OWNER = frozenset({OrganizationRole.OWNER})

# ── two representative targets ───────────────────────────────────────────────

ACCOUNTS = ImportTarget(
    key="probe.accounts",
    label_key="import.target.probe.accounts",
    permission="imports.manage",
    fields=(
        FieldSpec("email", "import.field.email", kind=FieldKind.EMAIL, required=True),
        FieldSpec("name", "import.field.name", required=True),
        FieldSpec("tier", "import.field.tier", options=("basic", "pro")),
        FieldSpec("seats", "import.field.seats", kind=FieldKind.INTEGER),
    ),
    match_keys=("email",),
    operations=(Operation.SKIP_DUPLICATE, Operation.CREATE),
    formats=(Format.CSV, Format.XLSX),
    max_rows=1000,
    version=2,
)

CONTACTS = ImportTarget(
    key="probe.contacts",
    label_key="import.target.probe.contacts",
    permission="imports.manage",
    fields=(
        FieldSpec("given_name", "import.field.given_name", required=True),
        FieldSpec("family_name", "import.field.family_name", required=True),
        FieldSpec("email", "import.field.email", kind=FieldKind.EMAIL),
        FieldSpec("phone", "import.field.phone"),
        FieldSpec("birth_date", "import.field.birth_date", kind=FieldKind.DATE),
        FieldSpec("newsletter", "import.field.newsletter", kind=FieldKind.BOOLEAN),
    ),
    match_keys=("email",),
    operations=(Operation.SKIP_DUPLICATE,),
    formats=(Format.CSV, Format.XLSX),
    max_rows=1000,
)

#: The same shape as `ACCOUNTS`, taking three rows a run.
SMALL = replace(ACCOUNTS, key="probe.small", max_rows=3)


def workbook(rows: list[list[object]], *, sheet: str = "Sheet1") -> bytes:
    from openpyxl import Workbook

    made = Workbook()
    active = made.active
    assert active is not None
    active.title = sheet
    for row in rows:
        active.append(row)
    buffer = io.BytesIO()
    made.save(buffer)
    return buffer.getvalue()


def as_csv(rows: list[list[object]]) -> bytes:
    return "".join(
        ",".join("" if cell is None else str(cell) for cell in row) + "\n" for row in rows
    ).encode()


#: Headers a customer's file actually arrives with, for each target: the
#: declared names, the names another system uses, a required field missing, a
#: column nothing matches, a name appearing twice, and a column with no name.
ACCOUNT_FILES: dict[str, list[list[object]]] = {
    "declared-names": [
        ["email", "name", "tier", "seats"],
        ["ada@example.com", "Analytical Engines Ltd", "pro", 12],
        ["grace@example.com", "Compilers Inc", "basic", 3],
    ],
    "another-systems-names": [
        ["E-Mail", " NAME ", "Tier", "Seats", "Created"],
        ["ada@example.com", "Analytical Engines Ltd", "pro", 12, date(2025, 3, 4)],
    ],
    "a-required-field-missing": [["name", "tier"], ["Analytical Engines Ltd", "pro"]],
    "unknown-columns": [
        ["email", "name", "favourite colour", "notes"],
        ["ada@example.com", "Ada", "green", "x"],
    ],
    "a-name-twice": [["email", "name", "name", None], ["a@x.io", "first", "second", "loose"]],
    "header-only": [["email", "name"]],
    "blank-rows-and-ragged": [
        ["email", "name", "tier"],
        [None, None, None],
        ["a@x.io"],
        ["b@x.io", "B", "pro", "extra"],
    ],
}

CONTACT_FILES: dict[str, list[list[object]]] = {
    "declared-names": [
        ["given_name", "family_name", "email", "phone", "birth_date", "newsletter"],
        ["Ada", "Lovelace", "ada@example.com", "+44 20 7946 0000", date(1815, 12, 10), True],
        ["Grace", "Hopper", "grace@example.com", None, date(1906, 12, 9), False],
    ],
    "another-systems-names": [
        ["Given Name", "FAMILY-NAME", "Email", "Phone", "Birth Date", "Newsletter?", "Id"],
        ["Ada", "Lovelace", "ada@example.com", 442079460000, date(1815, 12, 10), True, 1],
    ],
    "a-required-field-missing": [["given_name", "email"], ["Ada", "ada@example.com"]],
    "nothing-matches": [["Vorname", "Nachname"], ["Ada", "Lovelace"]],
}

CASES = [
    (target, name, rows)
    for target, files in ((ACCOUNTS, ACCOUNT_FILES), (CONTACTS, CONTACT_FILES))
    for name, rows in files.items()
]


# ── what the route answered until GR-352B ────────────────────────────────────


def analysed_as_before(raw: bytes, target: ImportTarget, fmt: Format) -> store.Analysis:
    """`core/imports.analyse` as it stood at `b24a72d`, kept as the reference.

    Not shared with the function under test, so the two cannot drift together:
    this one still calls the canonical readers.
    """
    limits = store.limits_for(target, rows=False)
    if fmt is Format.XLSX:
        read = read_workbook(raw, limit=store.PREVIEW, ceiling=target.max_rows, limits=limits)
        return store.Analysis(
            delimiter=",",
            encoding="utf-8",
            columns=read.header.columns,
            suggested=suggest(target, read.header.columns),
            preview=[row.cells for row in read.rows],
            rows_seen=read.rows_seen,
            over_ceiling=read.rows_seen > target.max_rows,
            replaced=False,
            format=fmt,
            sheet=read.sheet,
            identity=read.identity,
            template=compare(target, read.header.columns, found=read.identity),
        )
    preflight(raw, Format.CSV, limits)
    decoded = decode(raw)
    delimiter = sniff_delimiter(decoded.text)
    lines = decoded.text.splitlines(keepends=True)
    header = read_header(lines, delimiter=delimiter)
    rows = list(read_rows(lines, header=header, delimiter=delimiter, limit=store.PREVIEW))
    seen = count_rows(lines, delimiter=delimiter, ceiling=target.max_rows)
    return store.Analysis(
        delimiter=delimiter,
        encoding=decoded.encoding,
        columns=header.columns,
        suggested=suggest(target, header.columns),
        preview=[row.cells for row in rows],
        rows_seen=seen,
        over_ceiling=seen > target.max_rows,
        replaced=decoded.replaced,
        format=fmt,
        template=compare(target, header.columns),
    )


@pytest.mark.parametrize("fmt", [Format.CSV, Format.XLSX])
@pytest.mark.parametrize(
    ("target", "name", "rows"), CASES, ids=[f"{case[0].key}:{case[1]}" for case in CASES]
)
def test_the_analysis_is_the_one_the_route_gave_before(
    target: ImportTarget, name: str, rows: list[list[object]], fmt: Format
) -> None:
    raw = workbook(rows) if fmt is Format.XLSX else as_csv(rows)
    assert store.analyse(raw, target, fmt) == analysed_as_before(raw, target, fmt), name


def test_the_suggested_mapping_is_what_it_was_spelled_out_once() -> None:
    # The comparison above, with one answer written down, so a change to the
    # suggestion rule cannot move both sides of it together unnoticed.
    rows = ACCOUNT_FILES["another-systems-names"]
    for raw, fmt in ((workbook(rows), Format.XLSX), (as_csv(rows), Format.CSV)):
        found = store.analyse(raw, ACCOUNTS, fmt)
        assert found.columns == ("E-Mail", "NAME", "Tier", "Seats", "Created")
        assert found.suggested == {
            "E-Mail": "email",
            "NAME": "name",
            "Tier": "tier",
            "Seats": "seats",
        }
        assert found.template is not None
        assert found.template.verdict == "unknown_columns"
        assert found.template.unknown_columns == ("Created",)
        assert found.template.missing_required == ()

    missing = store.analyse(
        workbook(CONTACT_FILES["a-required-field-missing"]), CONTACTS, Format.XLSX
    )
    assert missing.template is not None
    assert missing.template.verdict == "incompatible"
    assert missing.template.missing_required == ("family_name",)


def test_a_template_the_target_gave_out_is_recognised_as_before() -> None:
    raw = render_xlsx(ACCOUNTS)
    found = store.analyse(raw, ACCOUNTS, Format.XLSX)
    assert found == analysed_as_before(raw, ACCOUNTS, Format.XLSX)
    assert found.sheet == "Data"
    assert found.template is not None and found.template.version_found == 2
    assert found.template.stale is False


def test_the_sample_is_still_the_first_two_hundred_rows() -> None:
    assert store.PREVIEW == 200
    rows: list[list[object]] = [["email", "name"]] + [[f"{n}@x.io", f"n{n}"] for n in range(350)]
    for raw, fmt in ((workbook(rows), Format.XLSX), (as_csv(rows), Format.CSV)):
        found = store.analyse(raw, ACCOUNTS, fmt)
        assert len(found.preview) == 200
        assert found.preview[0] == {"email": "0@x.io", "name": "n0"}
        assert found.preview[-1] == {"email": "199@x.io", "name": "n199"}
        assert found.rows_seen == 350 and found.over_ceiling is False


@pytest.mark.parametrize("fmt", [Format.CSV, Format.XLSX])
def test_a_file_over_the_ceiling_is_counted_one_past_it_and_no_further(fmt: Format) -> None:
    rows: list[list[object]] = [["email", "name"]] + [[f"{n}@x.io", f"n{n}"] for n in range(40)]
    raw = workbook(rows) if fmt is Format.XLSX else as_csv(rows)
    found = store.analyse(raw, SMALL, fmt)
    assert found == analysed_as_before(raw, SMALL, fmt)
    assert found.over_ceiling is True and found.rows_seen == SMALL.max_rows + 1


# ── the routes ───────────────────────────────────────────────────────────────


class _Result:
    def first(self) -> None:
        return None

    def __iter__(self) -> Iterator[Any]:
        return iter(())


class _Session:
    def __init__(self) -> None:
        self.commits = 0

    async def execute(self, *_: object, **__: object) -> _Result:
        return _Result()

    async def commit(self) -> None:
        self.commits += 1

    async def rollback(self) -> None:
        return None


def a_run(target: ImportTarget, fmt: Format) -> store.Run:
    return store.Run(
        id=RUN_ID,
        target=target.key,
        status=RunState.CREATED.value,
        format=fmt.value,
        source_file_id="22222222-2222-2222-2222-222222222222",
        delimiter=",",
        encoding="utf-8",
        columns=(),
        mapping={},
        operation=Operation.SKIP_DUPLICATE.value,
        rows_total=0,
        rows_valid=0,
        errors_total=0,
        errors_cut=False,
        error=None,
        requested_by=SUBJECT,
        committed_by=None,
        created_at=datetime.now(UTC),
        started_at=None,
        finished_at=None,
    )


class Routed:
    """The application with one run, one file, and nothing behind either."""

    def __init__(self, monkeypatch: pytest.MonkeyPatch) -> None:
        self.session = _Session()
        self.run = a_run(ACCOUNTS, Format.CSV)
        self.raw = b""
        self.mapped: list[dict[str, Any]] = []

        async def get(_session: object, _run_id: str) -> store.Run:
            return self.run

        async def source_bytes(*_: object) -> bytes:
            return self.raw

        async def set_mapping(_session: object, _run: store.Run, **recorded: Any) -> None:  # noqa: ANN401
            self.mapped.append(recorded)

        monkeypatch.setattr(store, "get", get)
        monkeypatch.setattr(store, "source_bytes", source_bytes)
        monkeypatch.setattr(store, "set_mapping", set_mapping)

        async def _db() -> AsyncIterator[_Session]:
            yield self.session

        app.dependency_overrides[require_auth] = lambda: JWTClaims(
            sub=SUBJECT, roles=OWNER, organization_id="org-1"
        )
        app.dependency_overrides[require_tenant] = lambda: TenantContext(
            id=TENANT, slug="alpha", name="Alpha", organization_id="org-1", user_id=SUBJECT
        )
        app.dependency_overrides[get_db] = _db
        app.dependency_overrides[tenant_storage] = lambda: SimpleNamespace(store=None)

    def holding(self, raw: bytes, target: ImportTarget, fmt: Format) -> None:
        self.raw = raw
        self.run = a_run(target, fmt)


@pytest.fixture
def routed(monkeypatch: pytest.MonkeyPatch) -> Iterator[Routed]:
    registry.clear()
    registry.extend([ACCOUNTS, CONTACTS, SMALL])
    yield Routed(monkeypatch)
    registry.clear()
    app.dependency_overrides.clear()


@pytest.fixture
def no_readers(monkeypatch: pytest.MonkeyPatch) -> None:
    """From here on, reaching either whole-file read fails the test."""
    import openpyxl
    from koras_import import reading, reading_xlsx

    def workbook_reached(*_: object, **__: object) -> None:
        raise AssertionError("a route reached openpyxl.load_workbook")

    def decode_reached(*_: object, **__: object) -> None:
        raise AssertionError("a route decoded the whole file")

    monkeypatch.setattr(openpyxl, "load_workbook", workbook_reached)
    monkeypatch.setattr(reading_xlsx, "read_workbook", workbook_reached)
    # The store itself held both names until GR-352C, when the dry run and the
    # commit stopped calling them as well; there is nothing left there to
    # replace.
    assert not hasattr(store, "read_workbook")
    assert not hasattr(store, "decode")
    monkeypatch.setattr(reading, "decode", decode_reached)


ANALYSIS = f"/api/v1/imports/{RUN_ID}/analysis"
MAPPING = f"/api/v1/imports/{RUN_ID}/mapping"

#: Every field the analysis route has answered since 2026-09-29. A change here
#: is a change to a public response, and GR-352B made none.
ANALYSIS_FIELDS = {
    "columns",
    "suggested",
    "preview",
    "rows_seen",
    "over_ceiling",
    "replaced",
    "format",
    "sheet",
    "template",
}


@pytest.mark.parametrize("fmt", [Format.CSV, Format.XLSX])
def test_the_analysis_route_answers_what_it_answered(routed: Routed, fmt: Format) -> None:
    rows = ACCOUNT_FILES["another-systems-names"]
    raw = workbook(rows) if fmt is Format.XLSX else as_csv(rows)
    routed.holding(raw, ACCOUNTS, fmt)
    answer = TestClient(app).get(ANALYSIS)
    assert answer.status_code == 200
    body = answer.json()
    assert set(body) == ANALYSIS_FIELDS
    before = analysed_as_before(raw, ACCOUNTS, fmt)
    assert body["columns"] == list(before.columns)
    assert body["suggested"] == before.suggested
    assert body["preview"] == before.preview
    assert body["rows_seen"] == before.rows_seen
    assert body["over_ceiling"] is before.over_ceiling
    assert body["replaced"] is before.replaced
    assert body["format"] == fmt.value
    assert body["sheet"] == before.sheet
    assert before.template is not None
    assert body["template"] == {
        "verdict": before.template.verdict,
        "missing_required": list(before.template.missing_required),
        "unknown_columns": list(before.template.unknown_columns),
        "version_found": before.template.version_found,
        "stale": before.template.stale,
    }


@pytest.mark.parametrize("fmt", [Format.CSV, Format.XLSX])
@pytest.mark.usefixtures("no_readers")
def test_neither_route_reaches_a_reader(routed: Routed, fmt: Format) -> None:
    rows = ACCOUNT_FILES["declared-names"]
    routed.holding(workbook(rows) if fmt is Format.XLSX else as_csv(rows), ACCOUNTS, fmt)
    client = TestClient(app)
    assert client.get(ANALYSIS).status_code == 200
    answer = client.put(MAPPING, json={"mapping": {"email": "email", "name": "name"}})
    assert answer.status_code == 200, answer.text


def test_the_guard_above_does_fail_when_a_reader_is_reached(
    routed: Routed, monkeypatch: pytest.MonkeyPatch
) -> None:
    # The route as it was: handed the old analysis, the same guard is reached.
    # A guard that cannot fire proves nothing about the route behind it.
    import openpyxl

    def reached(*_: object, **__: object) -> None:
        raise AssertionError("reached")

    monkeypatch.setattr(openpyxl, "load_workbook", reached)
    monkeypatch.setattr(store, "analyse", lambda raw, target, fmt, **_: analysed_as_before(
        raw, target, fmt
    ))
    routed.holding(workbook(ACCOUNT_FILES["declared-names"]), ACCOUNTS, Format.XLSX)
    answer = TestClient(app).get(ANALYSIS)
    # `read_workbook` turns what `openpyxl` raised into its own refusal.
    assert answer.status_code == 422
    assert answer.json()["detail"]["code"] == "import_file_unreadable"


@pytest.mark.parametrize("fmt", [Format.CSV, Format.XLSX])
def test_the_mapping_route_asks_for_the_header_and_no_rows(
    routed: Routed, monkeypatch: pytest.MonkeyPatch, fmt: Format
) -> None:
    rows = ACCOUNT_FILES["another-systems-names"]
    raw = workbook(rows) if fmt is Format.XLSX else as_csv(rows)
    routed.holding(raw, ACCOUNTS, fmt)
    asked: list[int] = []
    produced: list[store.Analysis] = []
    analyse = store.analyse

    def recorded(raw: bytes, target: ImportTarget, fmt: Format, *, sample: int) -> store.Analysis:
        asked.append(sample)
        produced.append(analyse(raw, target, fmt, sample=sample))
        return produced[-1]

    monkeypatch.setattr(store, "analyse", recorded)
    answer = TestClient(app).put(MAPPING, json={"mapping": {"E-Mail": "email", "NAME": "name"}})
    assert answer.status_code == 200, answer.text
    assert asked == [0]
    assert produced[0].preview == []
    # What the run is given is what the old analysis would have given it.
    before = analysed_as_before(raw, ACCOUNTS, fmt)
    assert routed.mapped == [
        {
            "delimiter": before.delimiter,
            "encoding": before.encoding,
            "columns": before.columns,
            "mapping": {"E-Mail": "email", "NAME": "name"},
        }
    ]
    assert routed.session.commits == 1


def test_the_mapping_route_still_refuses_what_it_refused(routed: Routed) -> None:
    client = TestClient(app)
    rows = ACCOUNT_FILES["declared-names"]
    routed.holding(as_csv(rows), ACCOUNTS, Format.CSV)

    undeclared = client.put(MAPPING, json={"mapping": {"email": "password_hash"}})
    assert undeclared.status_code == 422
    assert undeclared.json()["detail"]["code"] == "import_mapping_refused"

    absent = client.put(MAPPING, json={"mapping": {"no such column": "email"}})
    assert absent.status_code == 422
    assert absent.json()["detail"]["code"] == "import_mapping_refused"

    long: list[list[object]] = [["email", "name"]] + [[f"{n}@x.io", "n"] for n in range(9)]
    for raw, fmt in ((as_csv(long), Format.CSV), (workbook(long), Format.XLSX)):
        routed.holding(raw, SMALL, fmt)
        too_long = client.put(MAPPING, json={"mapping": {"email": "email", "name": "name"}})
        assert too_long.status_code == 422, fmt
        assert too_long.json()["detail"]["code"] == "import_too_many_rows"

    assert routed.mapped == []


@pytest.mark.parametrize("path", ["analysis", "mapping"])
def test_an_unsafe_file_is_still_refused_by_both_routes(
    routed: Routed, monkeypatch: pytest.MonkeyPatch, path: str
) -> None:
    monkeypatch.setattr(store, "SAFETY_LIMITS", SafetyLimits(max_decoded_string_bytes=50_000))
    raw = ("email,name\n" + "ada@example.com,Ada\n" * 3000 + "\U0001f600,x\n").encode()
    routed.holding(raw, ACCOUNTS, Format.CSV)
    client = TestClient(app)
    answer = (
        client.get(ANALYSIS)
        if path == "analysis"
        else client.put(MAPPING, json={"mapping": {"email": "email", "name": "name"}})
    )
    assert answer.status_code == 413
    assert answer.json()["detail"]["code"] == "import_file_too_large"
    assert routed.mapped == []


# ── the event loop ───────────────────────────────────────────────────────────

#: How long the stand-in inspection waits to be released before it gives up.
#: Reached only when the route is wrong, so it is the cost of a failure and
#: not of a pass.
GIVES_UP_AFTER = 5.0


@pytest.mark.parametrize("path", ["analysis", "mapping"])
@pytest.mark.asyncio
async def test_another_request_is_answered_while_a_file_is_inspected(
    routed: Routed, monkeypatch: pytest.MonkeyPatch, path: str
) -> None:
    rows = ACCOUNT_FILES["declared-names"]
    raw = as_csv(rows)
    routed.holding(raw, ACCOUNTS, Format.CSV)
    entered = threading.Event()
    release = threading.Event()
    analyse = store.analyse

    def held(raw: bytes, target: ImportTarget, fmt: Format, *, sample: int) -> store.Analysis:
        # An inspection that takes as long as the test says it does.
        entered.set()
        if not release.wait(GIVES_UP_AFTER):
            raise ReadRefused("the event loop was held for the whole inspection")
        return analyse(raw, target, fmt, sample=sample)

    monkeypatch.setattr(store, "analyse", held)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        inspecting = asyncio.create_task(
            client.get(ANALYSIS)
            if path == "analysis"
            else client.put(MAPPING, json={"mapping": {"email": "email", "name": "name"}})
        )
        # Until the inspection has started there is nothing to be held by.
        while not entered.is_set():
            assert not inspecting.done(), inspecting.result().text
            await asyncio.sleep(0.001)

        # The inspection is now in progress and will stay so until released
        # below. A second request, to the same application, on the same loop.
        other = await client.get("/api/v1/imports/targets")
        assert other.status_code == 200
        assert {found["key"] for found in other.json()} == {
            ACCOUNTS.key,
            CONTACTS.key,
            SMALL.key,
        }
        # It was answered *during* the inspection, not after it.
        assert not inspecting.done()
        assert not release.is_set()

        release.set()
        answer = await inspecting
    assert answer.status_code == 200, answer.text


@pytest.mark.asyncio
async def test_the_loop_test_above_does_fail_for_an_inline_inspection(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # The same stand-in, called the way the routes called `analyse` before:
    # inline, on the loop. Nothing else can run, so nothing releases it, and
    # it gives up -- which is the failure the test above would report.
    release = threading.Event()
    progressed: list[str] = []

    def held(patience: float) -> str:
        return "released" if release.wait(patience) else "gave up"

    async def another_request() -> None:
        progressed.append("answered")
        release.set()

    other = asyncio.create_task(another_request())
    assert held(0.2) == "gave up"
    assert progressed == []
    await other

    # And through a thread, which is what `analysed` does, it is released.
    release.clear()
    progressed.clear()
    other = asyncio.create_task(another_request())
    assert await asyncio.to_thread(held, GIVES_UP_AFTER) == "released"
    assert progressed == ["answered"]
    await other


@pytest.mark.asyncio
async def test_analysed_is_analyse_on_another_thread() -> None:
    seen: list[int] = []
    raw = as_csv(ACCOUNT_FILES["declared-names"])
    here = threading.get_ident()

    def where(*_: object, **__: object) -> store.Analysis:
        seen.append(threading.get_ident())
        return analysed_as_before(raw, ACCOUNTS, Format.CSV)

    original = store.analyse
    store.analyse = where  # type: ignore[assignment]
    try:
        found = await store.analysed(raw, ACCOUNTS, Format.CSV)
    finally:
        store.analyse = original  # type: ignore[assignment]
    assert seen and seen[0] != here
    assert found == original(raw, ACCOUNTS, Format.CSV)


# ── how many sources at once ─────────────────────────────────────────────────


async def _sources_held_at_once(
    routed: Routed, monkeypatch: pytest.MonkeyPatch, path: str, requests: int
) -> tuple[int, list[int]]:
    """Send `requests` together. The most sources in hand at once, and the answers.

    A source is in hand from the moment the route has fetched it until its
    inspection has returned -- which is how long the route keeps the bytes.
    """
    raw = as_csv(ACCOUNT_FILES["declared-names"])
    routed.holding(raw, ACCOUNTS, Format.CSV)
    lock = threading.Lock()
    release = threading.Event()
    held = 0
    most = 0
    analyse = store.analyse

    async def source_bytes(*_: object) -> bytes:
        nonlocal held, most
        with lock:
            held += 1
            most = max(most, held)
        return raw

    def slow(raw: bytes, target: ImportTarget, fmt: Format, *, sample: int) -> store.Analysis:
        nonlocal held
        release.wait(GIVES_UP_AFTER)
        try:
            return analyse(raw, target, fmt, sample=sample)
        finally:
            with lock:
                held -= 1

    monkeypatch.setattr(store, "source_bytes", source_bytes)
    monkeypatch.setattr(store, "analyse", slow)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        sent = [
            asyncio.create_task(
                client.get(ANALYSIS)
                if path == "analysis"
                else client.put(MAPPING, json={"mapping": {"email": "email", "name": "name"}})
            )
            for _ in range(requests)
        ]
        # Long enough for every request that is going to fetch to have fetched.
        for _ in range(30):
            await asyncio.sleep(0.01)
        release.set()
        answers = await asyncio.gather(*sent)
    return most, [answer.status_code for answer in answers]


@pytest.mark.parametrize("path", ["analysis", "mapping"])
@pytest.mark.asyncio
async def test_no_more_sources_are_held_at_once_than_there_are_slots(
    routed: Routed, monkeypatch: pytest.MonkeyPatch, path: str
) -> None:
    """IMPORT-GAP-019. An analysis is small; the source it is handed is not.

    Each of these requests fetches its file whole, up to 64 MiB of it, before
    its inspection starts. Nothing bounded how many did that at once in one
    512 MB process.
    """
    most, answers = await _sources_held_at_once(routed, monkeypatch, path, requests=7)
    assert most == store.ANALYSIS_SLOTS == 2
    # Nobody is refused for arriving while the slots are taken: they wait.
    assert answers == [200] * 7


@pytest.mark.asyncio
async def test_the_count_above_does_rise_without_the_slot(
    routed: Routed, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(store, "ANALYSIS_SLOTS", 7)
    most, answers = await _sources_held_at_once(routed, monkeypatch, "analysis", requests=7)
    assert most == 7 and answers == [200] * 7


@pytest.mark.asyncio
async def test_a_refused_file_gives_its_slot_back(
    routed: Routed, monkeypatch: pytest.MonkeyPatch
) -> None:
    # More refusals than there are slots, one after another, and then a file
    # that passes: a slot kept by a refusal would leave this waiting for ever.
    monkeypatch.setattr(store, "SAFETY_LIMITS", SafetyLimits(max_decoded_string_bytes=64))
    routed.holding(as_csv(ACCOUNT_FILES["declared-names"]), ACCOUNTS, Format.CSV)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        for _ in range(store.ANALYSIS_SLOTS + 2):
            assert (await client.get(ANALYSIS)).status_code == 413
        monkeypatch.setattr(store, "SAFETY_LIMITS", SafetyLimits())
        answer = await asyncio.wait_for(client.get(ANALYSIS), timeout=GIVES_UP_AFTER)
    assert answer.status_code == 200
