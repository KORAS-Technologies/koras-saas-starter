"""A workbook as rows, by the rules the CSV reader already keeps.

The point of this module is what it does *not* add: a second header rule, a
second row rule, a second idea of what a cell is. `header_from` and `row_from`
are the CSV reader's, so a spreadsheet saved as CSV and the same spreadsheet
saved as XLSX name their columns identically and produce the same cells. The
only thing a workbook has that a CSV does not is a typed cell, and `cell_text`
turns every one into the text the validator expects before anything sees it.

**Bounded before a sheet is opened, and since 2026-10-01 before `openpyxl` is
called at all.** `preflight` streams the workbook's parts and refuses one
whose decoded strings, cells or columns are outside the safety envelope --
`openpyxl` builds the whole shared-string table on load, so a check made after
it has already paid for what it refuses. GR-352, and `safety.py` says what
was measured. The directory checks that were here first still run: a file
claiming more than `DECOMPRESSED_CEILING` is refused without decompressing any
of it, a workbook carrying a macro project is refused, and one with no
workbook part is not a workbook.

**The worker does not call `read_workbook`, since GR-352C.** It stays as the
description of what a workbook's rows *are* -- `openpyxl`'s own reading of
them -- and `test_streaming.py` holds the reader the worker does use to it,
file by file. It is not that reader because of what `openpyxl` builds beside
the rows: every merged range, hyperlink and data validation of the sheet as an
object, the stylesheet and the workbook part as trees, and a walk of every
other worksheet to its end. The safety pass counts none of it, and a workbook
inside every limit was measured at 2.6 GiB here. A product that calls this
function on a customer's file inherits that; `open_rows` is the one to call.

**Cells are read with their cached values.** A formula never runs; it
contributes whatever the spreadsheet last calculated, or nothing. A cell whose
first character is a formula trigger is data here and is guarded on the way
out, where it always was.
"""

from __future__ import annotations

import io
import math
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any
from xml.etree import ElementTree  # noqa: S405 - bounded input, no entities resolved

from .compatibility import PROPERTY_NAME
from .reading import MAX_CELL, Header, ReadRefused, Row, header_from, row_from
from .safety import (
    MAX_UNCOMPRESSED_BYTES,
    PROVISIONAL_LIMITS,
    Preflight,
    SafetyLimits,
    preflight,
)
from .targets import Format

#: Summed from the zip directory before any entry is opened. ADR 0012 D12.
#: The number lives in `safety.py` now, with the rest of the envelope.
DECOMPRESSED_CEILING = MAX_UNCOMPRESSED_BYTES

#: The document-properties part is a few hundred bytes. One larger than this
#: is not read, so a hostile part cannot make the identity check expensive.
MAX_PROPERTY_BYTES = 64 * 1024

#: The sheet the template writes and the reader prefers. A workbook with no
#: sheet of that name is read from its first sheet, because most files a
#: customer has were not made from the template.
DATA_SHEET = "Data"

_WORKBOOK_PART = "xl/workbook.xml"
_MACRO_PART = "xl/vbaProject.bin"
_PROPERTIES_PART = "docProps/custom.xml"
_VT = "{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}"
_UNREADABLE = "the file could not be read as a workbook"


@dataclass(frozen=True)
class WorkbookRead:
    header: Header
    #: Up to `limit` rows, in order.
    rows: tuple[Row, ...]
    #: Rows seen while counting, stopping once past `ceiling`.
    rows_seen: int
    sheet: str
    #: The template identity the file carried, or None.
    identity: str | None
    #: What the safety pass counted before the workbook was opened.
    preflight: Preflight | None = None


def inspect(raw: bytes) -> zipfile.ZipFile:
    """Refuse what is not a safe workbook before opening any of it."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile as broken:
        raise ReadRefused("the file is not a workbook") from broken
    names = set(archive.namelist())
    if _WORKBOOK_PART not in names:
        raise ReadRefused("the file is not a workbook")
    if _MACRO_PART in names:
        raise ReadRefused("a workbook with macros is not read; save it as .xlsx first")
    total = sum(info.file_size for info in archive.infolist())
    if total > DECOMPRESSED_CEILING:
        raise ReadRefused(
            f"this workbook would expand to {total} bytes, over the "
            f"{DECOMPRESSED_CEILING} limit"
        )
    return archive


def template_identity(archive: zipfile.ZipFile) -> str | None:
    """The identity property, read without opening the workbook itself."""
    try:
        info = archive.getinfo(_PROPERTIES_PART)
    except KeyError:
        return None
    if info.file_size > MAX_PROPERTY_BYTES:
        return None
    part = archive.read(_PROPERTIES_PART)
    # A properties part never declares a document type or an entity. One that
    # does is not read at all, which closes entity expansion before the
    # parser is reached rather than trusting the parser to bound it.
    if b"<!DOCTYPE" in part or b"<!ENTITY" in part:
        return None
    try:
        root = ElementTree.fromstring(part)  # noqa: S314 - bounded, entity-free
    except ElementTree.ParseError:
        return None
    for prop in root:
        if prop.get("name") != PROPERTY_NAME:
            continue
        for child in prop:
            if child.tag == f"{_VT}lpwstr" and child.text:
                return child.text.strip()
    return None


def cell_text(value: Any) -> str:  # noqa: ANN401 - whatever openpyxl handed over
    """A typed cell as the text the validator would have read from a CSV.

    - `None` is an empty cell.
    - A boolean is `true` or `false`, which the validator accepts.
    - An integral float is written without its `.0`, because a spreadsheet
      stores every number as a float and `42.0` is not a whole number to the
      integer check.
    - A non-finite float is written as its name, so the decimal check refuses
      it rather than this module deciding for it.
    - A date is ISO. A datetime at midnight is its date; one with a time of
      day keeps the time, so a date column refuses it and says so.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return "nan" if math.isnan(value) else ("inf" if value > 0 else "-inf")
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    text = str(value)
    if len(text) > MAX_CELL:
        raise cell_too_long(len(text))
    return text


def cell_too_long(characters: int) -> ReadRefused:
    """The refusal for a cell that is a file. One sentence, for both readers."""
    return ReadRefused(
        f"a cell is {characters} characters, over the {MAX_CELL} limit; "
        "this is a file in a cell rather than a value"
    )


def _trimmed(values: Sequence[Any]) -> list[str]:
    """Cells as text, with the trailing empties a workbook pads rows with removed."""
    cells = [cell_text(value) for value in values]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def read_workbook(
    raw: bytes,
    *,
    limit: int,
    ceiling: int,
    limits: SafetyLimits = PROVISIONAL_LIMITS,
) -> WorkbookRead:
    """The header, up to `limit` rows, and a count that stops past `ceiling`.

    **The safety pass is inside this function rather than asked of its
    callers**, so there is no way to reach `openpyxl` through here without it.
    A caller chooses the envelope; it cannot choose to have none.
    """
    checked = preflight(raw, Format.XLSX, limits)

    from openpyxl import load_workbook

    archive = inspect(raw)
    found = template_identity(archive)
    archive.close()

    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as broken:  # noqa: BLE001 - openpyxl raises several types
        raise ReadRefused(_UNREADABLE) from broken
    try:
        # `Data`, or the first worksheet. Among worksheets only: a chart sheet
        # has no rows, and one named `Data` used to be chosen and then fail
        # with an error no caller answered.
        worksheets = book.worksheets
        sheet = next(
            (found for found in worksheets if found.title == DATA_SHEET),
            worksheets[0] if worksheets else None,
        )
        # **The sheet about to be read is the sheet the safety pass scanned,
        # or it is not read.** The pass resolves the package by `openpyxl`'s
        # own rules and scans the part those rules name; this is the other
        # half, asked of `openpyxl` itself after it has resolved the package
        # its own way. If the two ever disagree -- a rule missed, a release
        # that changes one -- the answer is a refusal rather than rows from a
        # part nothing counted. IMPORT-DEF-017.
        if sheet is None or getattr(sheet, "_worksheet_path", None) != checked.sheet_part:
            raise ReadRefused(_UNREADABLE)
        if hasattr(sheet, "reset_dimensions"):
            sheet.reset_dimensions()
        rows_iter = sheet.iter_rows(values_only=True)
        header: Header | None = None
        rows: list[Row] = []
        seen = 0
        number = 0
        try:
            for number, values in enumerate(rows_iter, start=1):
                cells = _trimmed(values)
                if header is None:
                    if not cells:
                        raise ReadRefused("the file's first row names no columns")
                    header = header_from(cells)
                    continue
                if not any(cells):
                    continue
                seen += 1
                if seen > ceiling:
                    break
                if len(rows) < limit:
                    rows.append(row_from(number, cells, header))
        except ReadRefused:
            raise
        except Exception as broken:  # noqa: BLE001 - whatever a malformed sheet raises
            # A cell naming a string the table lacks, a number that is not
            # one, a reference that is not a reference: `openpyxl` raises
            # each as it meets it, from inside the iteration, and none of
            # them is a `ReadRefused`. They reached the caller as they were
            # until 2026-10-01 -- a failed job with a traceback for a
            # sentence. A sheet that cannot be read is a refusal.
            raise ReadRefused(_UNREADABLE) from broken
        if header is None:
            raise ReadRefused("the file is empty")
        return WorkbookRead(
            header=header,
            rows=tuple(rows),
            rows_seen=seen,
            sheet=sheet.title,
            identity=found,
            preflight=checked,
        )
    finally:
        book.close()


__all__ = [
    "DATA_SHEET",
    "DECOMPRESSED_CEILING",
    "MAX_PROPERTY_BYTES",
    "WorkbookRead",
    "cell_text",
    "cell_too_long",
    "inspect",
    "read_workbook",
    "template_identity",
]
