"""A file to start from, rendered from the declaration the validator reads.

**There is no second description of the schema.** A template is the target's
`fields`, written out: the same names the mapping suggests from, the same
kinds the validator checks, the same options it accepts. A product that adds
a field gets a template with that field the next time anybody downloads one,
and nothing has to be kept level.

**Nothing here touches a tenant.** `render` takes a declaration and returns
bytes. No session, no locale, no request. A test can render every target of
every product without a database, and a template can never carry a customer's
row because it never sees one.

**Two formats.** CSV is a header row and nothing else, because a CSV has no
way to mark a row as an example that a header-parsing reader could trust, and
an example row somebody forgets to delete is an imported record. XLSX has a
Data sheet a person fills in and an Instructions sheet that explains every
column; the examples live there and in the header comments, which is where a
person filling the sheet looks. JSON is deferred (ADR 0012 D4) and refused by
name rather than by absence.

**Every string is guarded.** A field name, an option, an example or a help
text beginning with a character a spreadsheet reads as a formula is prefixed
with an apostrophe, the same rule the reporting writer applies. The values are
the product engineer's rather than a customer's; the guard is three lines; and
the file is opened by the one person a product least wants to surprise.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import TYPE_CHECKING

from .compatibility import PROPERTY_NAME, fingerprint, identity
from .targets import FieldKind, FieldSpec, Format, ImportTarget

if TYPE_CHECKING:
    from openpyxl.worksheet.datavalidation import DataValidation


class TemplateRefused(ValueError):
    """A template this engine cannot render, named."""


@dataclass(frozen=True)
class Rendered:
    content: bytes
    media_type: str
    extension: str
    filename: str


#: The formats a template can be rendered in. `Format.JSON` is in the
#: vocabulary and not here, by decision rather than by omission.
TEMPLATE_FORMATS: tuple[Format, ...] = (Format.CSV, Format.XLSX)

MEDIA_TYPES: dict[Format, tuple[str, str]] = {
    Format.CSV: ("text/csv; charset=utf-8", "csv"),
    Format.XLSX: (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
}

#: Data validation is applied to this many rows below the header, or the
#: target's row ceiling if that is smaller. A validation over a million rows
#: is a file that takes seconds to open for a column nobody will fill.
VALIDATED_ROWS = 10_000

#: A workbook list validation is limited to 255 characters inline. Longer
#: option sets go on a sheet of their own and are referenced by range.
_INLINE_LIST_LIMIT = 255

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

#: The one day every synthesised date example is: unremarkable, valid, and
#: obviously not anybody's real data.
_EXAMPLE_DAY = "2026-01-31"


def guard(text: str) -> str:
    """A string that opens as text in a spreadsheet, whatever it starts with."""
    return "'" + text if text.startswith(_FORMULA_PREFIXES) else text


def example_for(spec: FieldSpec) -> str:
    """The declared example, or one synthesised from the kind.

    Deterministic, so two renders agree and a test can pin them. Never
    anything that could be mistaken for a real record.
    """
    if spec.example:
        return spec.example
    if spec.options:
        return spec.options[0]
    return {
        FieldKind.TEXT: "Example",
        FieldKind.INTEGER: "42",
        FieldKind.DECIMAL: "19.99",
        FieldKind.BOOLEAN: "yes",
        FieldKind.DATE: _EXAMPLE_DAY,
        FieldKind.EMAIL: "name@example.com",
    }[spec.kind]


def format_hint(spec: FieldSpec) -> str:
    """One phrase saying how to write a value of this kind, or nothing."""
    if spec.options:
        return "one of: " + ", ".join(spec.options)
    return {
        FieldKind.TEXT: "",
        FieldKind.INTEGER: "whole number",
        FieldKind.DECIMAL: "number with a decimal point, no thousands separator",
        FieldKind.BOOLEAN: "yes or no",
        FieldKind.DATE: "YYYY-MM-DD",
        FieldKind.EMAIL: "email address",
    }[spec.kind]


def type_label(spec: FieldSpec) -> str:
    if spec.options:
        return "choice"
    return spec.kind.value


def filename(target: ImportTarget, fmt: Format) -> str:
    """`crm-accounts-template-v2.xlsx`: for a person, and trusted for nothing."""
    _, extension = MEDIA_TYPES[fmt]
    return f"{target.key.replace('.', '-')}-template-v{target.version}.{extension}"


def render(target: ImportTarget, fmt: Format) -> Rendered:
    """The template for a target, in a format it accepts and this can render."""
    if fmt not in TEMPLATE_FORMATS:
        raise TemplateRefused(f"a {fmt.value} template cannot be rendered")
    if not target.accepts(fmt):
        raise TemplateRefused(f"target {target.key} does not accept {fmt.value} files")
    content = render_csv(target) if fmt is Format.CSV else render_xlsx(target)
    media_type, extension = MEDIA_TYPES[fmt]
    return Rendered(
        content=content,
        media_type=media_type,
        extension=extension,
        filename=filename(target, fmt),
    )


def render_csv(target: ImportTarget) -> bytes:
    """The header row alone: UTF-8 with a byte-order mark, comma, CRLF.

    The mark is what makes a spreadsheet open the file with accented
    characters intact rather than guessing an encoding. CRLF because both
    major spreadsheets write it and every CSV reader accepts it.
    """
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([guard(spec.name) for spec in target.fields])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def render_xlsx(target: ImportTarget) -> bytes:
    """A Data sheet to fill in, an Instructions sheet to read, one property."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.packaging.custom import StringProperty
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    data = book.active
    if data is None:  # pragma: no cover - openpyxl always opens with one sheet
        data = book.create_sheet("Data")
    data.title = "Data"

    bold = Font(bold=True)
    required_fill = PatternFill(fill_type="solid", start_color="FFFCE8B2", end_color="FFFCE8B2")
    rows = min(target.max_rows, VALIDATED_ROWS)
    lists = None

    for index, spec in enumerate(target.fields, start=1):
        cell = data.cell(row=1, column=index, value=guard(spec.name))
        cell.font = bold
        if spec.required:
            cell.fill = required_fill
        note = _comment_for(spec)
        if note:
            cell.comment = Comment(note, "KORAS")
        letter = get_column_letter(index)
        width = max(len(spec.name), len(example_for(spec)), 12) + 2
        data.column_dimensions[letter].width = min(width, 60)

        validation = _validation_for(spec)
        if validation is None:
            continue
        if validation.type == "list" and validation.formula1 is None:
            # Too long for an inline list: put the options on their own
            # sheet and reference the range.
            if lists is None:
                lists = book.create_sheet("Lists")
            column = lists.max_column + (1 if lists.max_row > 1 or lists["A1"].value else 0)
            for offset, option in enumerate(spec.options, start=1):
                lists.cell(row=offset, column=column, value=guard(option))
            list_letter = get_column_letter(column)
            validation.formula1 = f"=Lists!${list_letter}$1:${list_letter}${len(spec.options)}"
        validation.add(f"{letter}2:{letter}{rows + 1}")
        data.add_data_validation(validation)

    data.freeze_panes = "A2"

    instructions = book.create_sheet("Instructions")
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 12
    instructions.column_dimensions["C"].width = 12
    instructions.column_dimensions["D"].width = 40
    instructions.column_dimensions["E"].width = 32
    instructions.column_dimensions["F"].width = 24
    instructions.column_dimensions["G"].width = 60

    heading = [
        ("Import template for", guard(target.key)),
        ("Version", f"v{target.version}"),
        ("Fingerprint", fingerprint(target)),
        ("Rendered", datetime.now(UTC).strftime("%Y-%m-%d")),
        ("Rows per import", str(target.max_rows)),
        (
            "Fill in",
            'The "Data" sheet. Keep row 1 as it is: the importer reads the column names.',
        ),
        ("Dates", "Write dates as YYYY-MM-DD."),
        ("Numbers", "Write decimals with a point and no thousands separator."),
        ("Yes or no", "Write yes or no."),
    ]
    for label, value in heading:
        instructions.append([label, value])
        instructions.cell(row=instructions.max_row, column=1).font = bold
    instructions.append([])
    instructions.append(
        ["Column", "Required", "Type", "Format", "Allowed values", "Example", "Help"]
    )
    for index in range(1, 8):
        instructions.cell(row=instructions.max_row, column=index).font = bold
    for spec in target.fields:
        instructions.append(
            [
                guard(spec.name),
                "required" if spec.required else "optional",
                type_label(spec),
                guard(format_hint(spec)),
                guard(", ".join(spec.options)),
                guard(example_for(spec)),
                guard(spec.help),
            ]
        )
    for row in instructions.iter_rows(min_row=1, max_row=instructions.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # The stubs shipped for openpyxl predate custom document properties.
    book.custom_doc_props.append(  # type: ignore[attr-defined]
        StringProperty(name=PROPERTY_NAME, value=identity(target))
    )

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _comment_for(spec: FieldSpec) -> str:
    parts = []
    if spec.help:
        parts.append(spec.help)
    hint = format_hint(spec)
    if hint:
        parts.append(hint.capitalize() + ".")
    parts.append(
        ("Required. " if spec.required else "Optional. ") + f"Example: {example_for(spec)}"
    )
    return guard(" ".join(parts))


def _validation_for(spec: FieldSpec) -> DataValidation | None:
    from openpyxl.worksheet.datavalidation import DataValidation

    if spec.options:
        joined = ",".join(spec.options)
        inline = f'"{joined}"' if len(joined) <= _INLINE_LIST_LIMIT else None
        return DataValidation(
            type="list",
            formula1=inline,
            allow_blank=not spec.required,
            showErrorMessage=True,
            errorTitle="Not an option",
            error="Choose one of the listed values.",
        )
    if spec.kind is FieldKind.DATE:
        return DataValidation(
            type="date",
            operator="greaterThan",
            formula1="1",
            allow_blank=not spec.required,
            showErrorMessage=True,
            errorTitle="Not a date",
            error="Write a date as YYYY-MM-DD.",
        )
    if spec.kind is FieldKind.INTEGER:
        return DataValidation(
            type="whole",
            operator="between",
            formula1="-999999999999",
            formula2="999999999999",
            allow_blank=not spec.required,
            showErrorMessage=True,
            errorTitle="Not a whole number",
            error="Write a whole number.",
        )
    if spec.kind is FieldKind.DECIMAL:
        return DataValidation(
            type="decimal",
            operator="between",
            formula1="-999999999999",
            formula2="999999999999",
            allow_blank=not spec.required,
            showErrorMessage=True,
            errorTitle="Not a number",
            error="Write a number with a decimal point.",
        )
    return None


__all__ = [
    "MEDIA_TYPES",
    "TEMPLATE_FORMATS",
    "VALIDATED_ROWS",
    "Rendered",
    "TemplateRefused",
    "example_for",
    "filename",
    "format_hint",
    "guard",
    "render",
    "render_csv",
    "render_xlsx",
    "type_label",
]
