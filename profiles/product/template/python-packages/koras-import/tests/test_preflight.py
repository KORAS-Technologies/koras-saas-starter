"""The safety pass: what it counts, what it refuses, and that it refuses first.

GR-352. Three kinds of assertion, and the third is the one the slice exists
for.

**What a string costs.** The width of a string is set by its widest character,
so the arithmetic is tested against `sys.getsizeof` rather than against
itself.

**What is refused.** Each dimension has a workbook and a CSV that break it and
nothing else, with limits passed in small so no test builds a large file.

**That refusal comes before the reader.** A refused workbook never reaches
`openpyxl.load_workbook`, proved by replacing it with something that fails the
test if called; and the scan's own memory is measured while it reads far more
text than it is allowed to keep.

Workbooks here are written as XML by hand where the shape matters -- inline
strings, rich text, a renamed string table, a lying `<dimension>` -- because
`openpyxl` writes only one of the shapes a customer's spreadsheet may arrive
in, and a scan that was only ever shown that one would be safe only for it.
"""

from __future__ import annotations

import io
import sys
import tracemalloc
import zipfile
from collections.abc import Sequence
from dataclasses import replace
from typing import Any
from xml.sax.saxutils import escape

import pytest
from koras_import import (
    ENVELOPE_CODES,
    PROVISIONAL_LIMITS,
    Format,
    PreflightRefused,
    ReadRefused,
    RefusalCode,
    SafetyLimits,
    count_rows,
    decode,
    decoded_cost,
    preflight,
    read_workbook,
    safety,
    sniff_delimiter,
    string_cost,
    survey,
    width_of,
)

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE = "http://schemas.openxmlformats.org/package/2006/relationships"
EMOJI = "\U0001f600"


# ── building workbooks by hand ───────────────────────────────────────────────


def package(parts: dict[str, str | bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    return buffer.getvalue()


def sst(strings: list[str]) -> str:
    items = "".join(f"<si><t>{escape(text)}</t></si>" for text in strings)
    return f'<sst xmlns="{MAIN}">{items}</sst>'


def worksheet(rows: str, *, dimension: str = "") -> str:
    declared = f'<dimension ref="{dimension}"/>' if dimension else ""
    return f'<worksheet xmlns="{MAIN}">{declared}<sheetData>{rows}</sheetData></worksheet>'


def letters(column: int) -> str:
    name = ""
    while column:
        column, rest = divmod(column - 1, 26)
        name = chr(65 + rest) + name
    return name


def shared_rows(grid: list[list[int | None]]) -> str:
    """Rows of shared-string references; None is a cell left out."""
    out = []
    for number, row in enumerate(grid, start=1):
        cells = "".join(
            f'<c r="{letters(column)}{number}" t="s"><v>{index}</v></c>'
            for column, index in enumerate(row, start=1)
            if index is not None
        )
        out.append(f'<row r="{number}">{cells}</row>')
    return "".join(out)


def inline_rows(grid: list[list[str]]) -> str:
    out = []
    for number, row in enumerate(grid, start=1):
        cells = "".join(
            f'<c r="{letters(column)}{number}" t="inlineStr"><is><t>{escape(text)}</t></is></c>'
            for column, text in enumerate(row, start=1)
        )
        out.append(f'<row r="{number}">{cells}</row>')
    return "".join(out)


def book(
    sheet: str,
    strings: str | None = None,
    *,
    sheets: dict[str, str] | None = None,
    strings_part: str = "xl/sharedStrings.xml",
    extra: dict[str, str | bytes] | None = None,
) -> bytes:
    """A workbook `openpyxl` will open, around the XML a test cares about."""
    named = sheets or {"Sheet1": sheet}
    overrides = [
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    ]
    entries = []
    relations = []
    parts: dict[str, str | bytes] = {}
    for position, (name, body) in enumerate(named.items(), start=1):
        overrides.append(
            f'<Override PartName="/xl/worksheets/sheet{position}.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        )
        entries.append(f'<sheet name="{name}" sheetId="{position}" r:id="rId{position}"/>')
        relations.append(
            f'<Relationship Id="rId{position}" Type="{RELS}/worksheet" '
            f'Target="worksheets/sheet{position}.xml"/>'
        )
        parts[f"xl/worksheets/sheet{position}.xml"] = body
    if strings is not None:
        overrides.append(
            f'<Override PartName="/{strings_part}" ContentType="application/vnd.'
            'openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        )
        relations.append(
            f'<Relationship Id="rIdS" Type="{RELS}/sharedStrings" '
            f'Target="{strings_part.removeprefix("xl/")}"/>'
        )
        parts[strings_part] = strings
    parts["[Content_Types].xml"] = (
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
        'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
        + "".join(overrides)
        + "</Types>"
    )
    parts["_rels/.rels"] = (
        f'<Relationships xmlns="{PACKAGE}"><Relationship Id="rId1" '
        f'Type="{RELS}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
    )
    parts["xl/workbook.xml"] = (
        f'<workbook xmlns="{MAIN}" xmlns:r="{RELS}"><sheets>{"".join(entries)}</sheets></workbook>'
    )
    parts["xl/_rels/workbook.xml.rels"] = (
        f'<Relationships xmlns="{PACKAGE}">{"".join(relations)}</Relationships>'
    )
    parts.update(extra or {})
    return package(parts)


def written(rows: Sequence[Sequence[object]], *, sheet: str = "Sheet1") -> bytes:
    """A workbook as `openpyxl` itself writes one."""
    from openpyxl import Workbook

    made = Workbook()
    active = made.active
    assert active is not None
    active.title = sheet
    for row in rows:
        active.append(list(row))
    buffer = io.BytesIO()
    made.save(buffer)
    return buffer.getvalue()


def limits(**changed: Any) -> SafetyLimits:  # noqa: ANN401 - whichever limit a test moves
    return replace(PROVISIONAL_LIMITS, **changed)


def refused(raw: bytes, fmt: Format, envelope: SafetyLimits, **options: str) -> PreflightRefused:
    with pytest.raises(PreflightRefused) as caught:
        preflight(raw, fmt, envelope, **options)
    return caught.value


def no_openpyxl(monkeypatch: pytest.MonkeyPatch) -> None:
    """From here on, reaching `openpyxl.load_workbook` fails the test."""
    import openpyxl

    def reached(*_: object, **__: object) -> None:
        raise AssertionError("openpyxl was reached with a file the safety pass refuses")

    monkeypatch.setattr(openpyxl, "load_workbook", reached)


# ── what a string costs ──────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("text", "width"),
    [
        ("", 1),
        ("plain ascii", 1),
        ("café über", 1),  # Latin-1: still one byte a character
        ("złoty €", 2),  # BMP: Polish, the euro sign
        ("漢字", 2),  # CJK
        (EMOJI, 4),  # astral
        ("a" * 1000 + EMOJI, 4),  # mostly ASCII, one astral character
    ],
)
def test_the_width_is_the_widest_character_and_the_cost_bounds_the_real_string(
    text: str, width: int
) -> None:
    assert width_of(text) == width
    assert string_cost(text) == decoded_cost(len(text), width)
    # Against the interpreter rather than against the formula: the estimate is
    # only useful if it is never below what the string really takes.
    assert string_cost(text) >= sys.getsizeof(text)


def test_one_astral_character_quadruples_the_cost_of_everything_beside_it() -> None:
    plain = "a" * 10_000
    assert string_cost(plain + EMOJI) > 3.9 * string_cost(plain)
    # Counting code points alone -- what this replaces -- sees no difference.
    assert len(plain + EMOJI) == len(plain) + 1


def test_limits_refuse_a_nonsense_envelope() -> None:
    with pytest.raises(ValueError, match="max_cells"):
        limits(max_cells=0)
    with pytest.raises(ValueError, match="max_rows"):
        limits(max_rows=0)


# ── XLSX: what passes ────────────────────────────────────────────────────────


def test_a_small_workbook_passes_and_is_counted() -> None:
    raw = written([["Name", "Email"], ["Ada", "ada@example.com"], ["Bo", "bo@example.com"]])
    found = preflight(raw, Format.XLSX)
    assert found.safe
    assert (found.rows, found.columns, found.cells) == (2, 2, 6)
    assert found.source_bytes == len(raw)
    assert found.uncompressed_bytes is not None and found.uncompressed_bytes > len(raw)
    assert found.widest_character == 1
    assert found.decoded_string_bytes > 0


def test_a_realistic_workbook_passes_and_agrees_with_the_reader_about_rows() -> None:
    header: list[object] = [f"Column {index}" for index in range(12)]
    rows: list[list[object]] = [header]
    for number in range(2000):
        rows.append([f"value {number}-{index}" for index in range(10)] + [number, number / 3])
    raw = written(rows)
    found = preflight(raw, Format.XLSX)
    read = read_workbook(raw, limit=10, ceiling=50_000)
    assert found.rows == read.rows_seen == 2000
    assert found.columns == 12
    assert found.cells == 2001 * 12
    assert read.preflight == found


def test_typed_cells_are_counted_and_never_reinterpreted() -> None:
    from datetime import date

    raw = written([["When", "Paid", "Total"], [date(2026, 1, 31), True, 12.5]])
    read = read_workbook(raw, limit=10, ceiling=10)
    assert read.rows[0].cells == {"When": "2026-01-31", "Paid": "true", "Total": "12.5"}
    assert read.preflight is not None and read.preflight.rows == 1


def test_rows_are_counted_by_the_readers_rule() -> None:
    # Header, a row, a row of nothing, a row whose only cell is an empty
    # string, a row. The reader calls that two rows.
    strings = sst(["Name", "Ada", "", "Bo"])
    raw = book(
        worksheet(
            shared_rows([[0], [1]])
            + '<row r="3"></row>'
            + '<row r="4"><c r="A4" t="s"><v>2</v></c></row>'
            + '<row r="5"><c r="A5" t="s"><v>3</v></c></row>'
        ),
        strings,
    )
    assert preflight(raw, Format.XLSX).rows == 2
    assert read_workbook(raw, limit=10, ceiling=10).rows_seen == 2


def test_only_the_sheet_the_reader_reads_is_counted() -> None:
    wide = inline_rows([["x"] * 40 for _ in range(30)])
    narrow = inline_rows([["Name"], ["Ada"]])
    envelope = limits(max_cells=100, max_columns=10)
    # The first sheet is read when none is called Data.
    assert preflight(
        book("", sheets={"First": worksheet(narrow), "Lookup": worksheet(wide)}),
        Format.XLSX,
        envelope,
    ).cells == 2
    # `Data` is read wherever it sits, so a wide one is refused...
    assert (
        refused(
            book("", sheets={"First": worksheet(narrow), "Data": worksheet(wide)}),
            Format.XLSX,
            envelope,
        ).code
        is RefusalCode.TOO_MANY_COLUMNS
    )
    # ...and a narrow one is not, whatever sits before it.
    assert preflight(
        book("", sheets={"Lookup": worksheet(wide), "Data": worksheet(narrow)}),
        Format.XLSX,
        envelope,
    ).cells == 2


# ── XLSX: what is refused ────────────────────────────────────────────────────


def test_a_workbook_over_the_source_ceiling_is_refused_unopened() -> None:
    raw = written([["Name"], ["Ada"]])
    refusal = refused(raw, Format.XLSX, limits(max_source_bytes=len(raw) - 1))
    assert refusal.code is RefusalCode.SOURCE_TOO_LARGE
    assert (refusal.limit, refusal.observed) == (len(raw) - 1, len(raw))


def test_a_workbook_that_would_expand_past_the_ceiling_is_refused() -> None:
    raw = book(worksheet(shared_rows([[0]])), sst(["a" * 200_000]))
    refusal = refused(raw, Format.XLSX, limits(max_uncompressed_bytes=100_000))
    assert refusal.code is RefusalCode.UNCOMPRESSED_TOO_LARGE
    assert len(raw) < 100_000  # compressed, it was nowhere near


def test_shared_strings_over_the_decoded_budget_are_refused() -> None:
    raw = book(worksheet(shared_rows([[0, 1]])), sst(["a" * 30_000, "b" * 30_000]))
    assert preflight(raw, Format.XLSX, limits(max_decoded_string_bytes=70_000)).safe
    refusal = refused(raw, Format.XLSX, limits(max_decoded_string_bytes=50_000))
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE


def test_inline_strings_over_the_decoded_budget_are_refused() -> None:
    raw = book(worksheet(inline_rows([["a" * 30_000, "b" * 30_000]])))
    assert preflight(raw, Format.XLSX, limits(max_decoded_string_bytes=70_000)).safe
    refusal = refused(raw, Format.XLSX, limits(max_decoded_string_bytes=50_000))
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE


@pytest.mark.parametrize("shape", ["shared", "inline"])
def test_one_astral_character_in_an_ascii_string_is_costed_at_four_bytes(shape: str) -> None:
    def made(text: str) -> bytes:
        if shape == "shared":
            return book(worksheet(shared_rows([[0]])), sst([text]))
        return book(worksheet(inline_rows([[text]])))

    envelope = limits(max_decoded_string_bytes=60_000)
    plain = preflight(made("a" * 30_000), Format.XLSX, envelope)
    assert plain.safe and plain.widest_character == 1
    # The same string, one character longer.
    refusal = refused(made("a" * 30_000 + EMOJI), Format.XLSX, envelope)
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE
    wide = survey(made("a" * 30_000 + EMOJI), Format.XLSX, envelope)
    assert wide.widest_character == 4


def test_an_astral_character_written_as_a_reference_is_still_astral() -> None:
    # The file's bytes are pure ASCII; the string the reader builds is not.
    strings = f'<sst xmlns="{MAIN}"><si><t>{"a" * 30_000}&#x1F600;</t></si></sst>'
    raw = book(worksheet(shared_rows([[0]])), strings)
    assert strings.isascii()
    refusal = refused(raw, Format.XLSX, limits(max_decoded_string_bytes=60_000))
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE


def test_rich_text_runs_are_one_logical_string_at_its_widest_run() -> None:
    # Two ASCII runs and one astral run: one string, four bytes a character.
    runs = (
        f"<si><r><t>{'a' * 10_000}</t></r><r><rPr><b/></rPr><t>{'b' * 10_000}</t></r>"
        f"<r><t>{EMOJI}</t></r></si>"
    )
    raw = book(worksheet(shared_rows([[0]])), f'<sst xmlns="{MAIN}">{runs}</sst>')
    found = preflight(raw, Format.XLSX)
    assert found.decoded_string_bytes >= decoded_cost(20_001, 4)


def test_a_string_table_is_found_by_what_it_is_and_not_by_its_name() -> None:
    raw = book(
        worksheet(shared_rows([[0]])),
        sst(["a" * 60_000]),
        strings_part="xl/elsewhere/words.xml",
    )
    refusal = refused(raw, Format.XLSX, limits(max_decoded_string_bytes=50_000))
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE


def test_a_formula_is_costed_like_any_other_text_the_parser_holds() -> None:
    cell = f'<row r="1"><c r="A1"><f>{"A2+" * 20_000}1</f><v>1</v></c></row>'
    refusal = refused(
        book(worksheet(cell)), Format.XLSX, limits(max_decoded_string_bytes=50_000)
    )
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE


def test_cells_are_counted_as_they_are_and_refused_past_the_budget() -> None:
    raw = book(worksheet(inline_rows([["x"] * 10 for _ in range(10)])))
    assert preflight(raw, Format.XLSX, limits(max_cells=100)).cells == 100
    refusal = refused(raw, Format.XLSX, limits(max_cells=99))
    assert refusal.code is RefusalCode.TOO_MANY_CELLS
    assert refusal.observed == 100


def test_columns_come_from_the_cells_and_not_from_the_declared_dimension() -> None:
    # One cell, far to the right, in a sheet that declares itself two wide.
    far = '<row r="1"><c r="A1" t="inlineStr"><is><t>a</t></is></c>' + (
        '<c r="ZZ1" t="inlineStr"><is><t>b</t></is></c></row>'
    )
    raw = book(worksheet(far, dimension="A1:B1"))
    refusal = refused(raw, Format.XLSX, limits(max_columns=256))
    assert refusal.code is RefusalCode.TOO_MANY_COLUMNS
    assert refusal.observed == 702  # ZZ
    assert preflight(raw, Format.XLSX, limits(max_columns=702)).columns == 702


def test_cells_without_references_are_counted_by_position() -> None:
    row = "<row>" + "<c><v>1</v></c>" * 12 + "</row>"
    assert preflight(book(worksheet(row)), Format.XLSX).columns == 12


def test_rows_are_refused_only_when_the_caller_asks_for_a_ceiling() -> None:
    raw = written([["Name"]] + [[f"row {number}"] for number in range(5)])
    assert preflight(raw, Format.XLSX).rows == 5  # no ceiling: counted, not refused
    assert preflight(raw, Format.XLSX, limits(max_rows=5)).safe  # the header is not a row
    refusal = refused(raw, Format.XLSX, limits(max_rows=4))
    assert refusal.code is RefusalCode.TOO_MANY_ROWS


@pytest.mark.parametrize(
    "raw",
    [
        b"",
        b"Name,Email\nAda,ada@example.com\n",
        package({"hello.txt": "not a workbook"}),
    ],
)
def test_what_is_not_a_workbook_is_refused_as_such(raw: bytes) -> None:
    refusal = refused(raw, Format.XLSX, PROVISIONAL_LIMITS)
    assert refusal.code is RefusalCode.MALFORMED
    assert str(refusal) == "the file is not a workbook"


def test_a_broken_sheet_is_a_refusal_and_not_a_parser_error() -> None:
    raw = book(f'<worksheet xmlns="{MAIN}"><sheetData><row><c>')
    refusal = refused(raw, Format.XLSX, PROVISIONAL_LIMITS)
    assert refusal.code is RefusalCode.MALFORMED
    # Nothing of the parser in the sentence a person reads.
    assert "expat" not in str(refusal) and "line" not in str(refusal)


def test_a_part_declaring_a_document_type_refuses_the_workbook() -> None:
    bomb = (
        '<?xml version="1.0"?><!DOCTYPE sst [<!ENTITY a "aaaaaaaaaa">]>'
        f'<sst xmlns="{MAIN}"><si><t>&a;</t></si></sst>'
    )
    refusal = refused(book(worksheet(shared_rows([[0]])), bomb), Format.XLSX, PROVISIONAL_LIMITS)
    assert refusal.code is RefusalCode.MALFORMED


def test_a_token_that_never_ends_is_refused_before_it_becomes_a_string() -> None:
    # One attribute, several megabytes long. `expat` holds an attribute until
    # its tag closes, so this is the one shape that streaming does not bound.
    huge = f'<worksheet xmlns="{MAIN}" junk="{"a" * 8_000_000}"><sheetData/></worksheet>'
    raw = book(huge)
    tracemalloc.start()
    try:
        refusal = refused(raw, Format.XLSX, PROVISIONAL_LIMITS)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert refusal.code is RefusalCode.MALFORMED
    assert peak < 4_000_000


def test_a_workbook_with_macros_is_refused() -> None:
    raw = book(worksheet(""), extra={"xl/vbaProject.bin": b"\x00"})
    assert refused(raw, Format.XLSX, PROVISIONAL_LIMITS).code is RefusalCode.MACROS


def test_an_archive_with_far_too_many_entries_is_refused_before_it_is_opened(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    raw = package({f"junk/{number}": "" for number in range(50)})

    def reached(*_: object, **__: object) -> None:
        raise AssertionError("the archive was opened before its entries were counted")

    monkeypatch.setattr(zipfile, "ZipFile", reached)
    refusal = refused(raw, Format.XLSX, limits(max_archive_entries=49))
    assert refusal.code is RefusalCode.TOO_MANY_ENTRIES
    assert refusal.observed == 50


def test_every_envelope_refusal_is_a_read_refusal_the_callers_already_answer() -> None:
    raw = written([["Name"], ["Ada"]])
    with pytest.raises(ReadRefused):
        preflight(raw, Format.XLSX, limits(max_cells=1))
    assert RefusalCode.TOO_MANY_CELLS in ENVELOPE_CODES
    assert RefusalCode.TOO_MANY_ROWS not in ENVELOPE_CODES
    assert RefusalCode.MALFORMED not in ENVELOPE_CODES


def test_survey_reports_a_refusal_instead_of_raising_it() -> None:
    raw = written([["Name"], ["Ada"]])
    found = survey(raw, Format.XLSX, limits(max_cells=1))
    assert not found.safe
    assert found.refusal is not None and found.refusal.code is RefusalCode.TOO_MANY_CELLS
    assert found.source_bytes == len(raw)


# ── XLSX: refused before the reader ──────────────────────────────────────────


@pytest.mark.parametrize(
    ("raw", "envelope", "code"),
    [
        (
            book(worksheet(shared_rows([[0]])), sst(["a" * 30_000 + EMOJI])),
            {"max_decoded_string_bytes": 60_000},
            RefusalCode.DECODED_TOO_LARGE,
        ),
        (
            book(worksheet(inline_rows([["a" * 30_000 + EMOJI]]))),
            {"max_decoded_string_bytes": 60_000},
            RefusalCode.DECODED_TOO_LARGE,
        ),
        (
            book(worksheet(inline_rows([["x"] * 10 for _ in range(10)]))),
            {"max_cells": 50},
            RefusalCode.TOO_MANY_CELLS,
        ),
        (
            book(worksheet(inline_rows([["x"] * 10]))),
            {"max_columns": 5},
            RefusalCode.TOO_MANY_COLUMNS,
        ),
        (
            book(worksheet(inline_rows([["h"], ["a"], ["b"], ["c"]]))),
            {"max_rows": 2},
            RefusalCode.TOO_MANY_ROWS,
        ),
        (
            book(worksheet(shared_rows([[0]])), sst(["a" * 200_000])),
            {"max_uncompressed_bytes": 100_000},
            RefusalCode.UNCOMPRESSED_TOO_LARGE,
        ),
    ],
    ids=["astral-shared", "astral-inline", "cells", "columns", "rows", "uncompressed"],
)
def test_the_reader_never_opens_a_workbook_the_safety_pass_refuses(
    monkeypatch: pytest.MonkeyPatch,
    raw: bytes,
    envelope: dict[str, int],
    code: RefusalCode,
) -> None:
    no_openpyxl(monkeypatch)
    with pytest.raises(PreflightRefused) as caught:
        read_workbook(raw, limit=10, ceiling=10, limits=limits(**envelope))
    assert caught.value.code is code


def test_the_guard_above_does_fail_when_openpyxl_is_reached(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # The proof that the proof works: with a workbook that passes, the same
    # replacement is reached, and `read_workbook` turns it into its own
    # refusal. A guard that could not fire would make the test above vacuous.
    no_openpyxl(monkeypatch)
    with pytest.raises(ReadRefused, match="could not be read as a workbook") as caught:
        read_workbook(written([["Name"], ["Ada"]]), limit=10, ceiling=10)
    assert isinstance(caught.value.__cause__, AssertionError)


def test_the_scan_holds_none_of_the_text_it_reads() -> None:
    # Eight megabytes of shared strings, each with an astral character: about
    # 32 MiB once decoded the way the reader decodes it. The scan is allowed
    # all of it and must keep none of it.
    strings = sst([f"{'a' * 8_000}{EMOJI}{number}" for number in range(1000)])
    raw = book(worksheet(shared_rows([[number] for number in range(1000)])), strings)
    envelope = limits(max_decoded_string_bytes=1 << 30)
    tracemalloc.start()
    try:
        found = preflight(raw, Format.XLSX, envelope)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert found.decoded_string_bytes > 30 * 1024 * 1024
    assert peak < 4 * 1024 * 1024


# ── CSV: what passes ─────────────────────────────────────────────────────────


def reader_rows(raw: bytes) -> int:
    """What the reader itself calls the row count of this file."""
    text = decode(raw).text
    return count_rows(
        text.splitlines(keepends=True), delimiter=sniff_delimiter(text), ceiling=10**9
    )


def test_a_small_csv_passes_and_is_counted() -> None:
    raw = b"Name,Email\nAda,ada@example.com\n\nBo,bo@example.com\n"
    found = preflight(raw, Format.CSV)
    assert found.safe
    assert (found.rows, found.columns, found.cells) == (2, 2, 6)
    assert found.encoding == "utf-8-sig"
    assert found.uncompressed_bytes is None
    assert found.decoded_string_bytes == decoded_cost(len(raw), 1)


def test_a_realistic_csv_passes_and_agrees_with_the_reader_about_rows() -> None:
    lines = ["Name;City;Note"] + [
        f'Person {number};"Köln; Ehrenfeld";note {number}\r\n' for number in range(5000)
    ]
    raw = ("\r\n".join(lines[:1]) + "\r\n" + "".join(lines[1:])).encode("cp1252")
    found = preflight(raw, Format.CSV)
    assert found.safe and found.encoding == "cp1252"
    assert found.rows == reader_rows(raw) == 5000
    assert found.columns == 3


def test_chunk_boundaries_change_nothing(monkeypatch: pytest.MonkeyPatch) -> None:
    # Every awkward thing a boundary can fall inside: a multi-byte character,
    # a `\r\n`, a quoted field with a line break in it.
    body = "".join(
        f'{number},"multi\nline {EMOJI}",złoty\r\n' for number in range(300)
    )
    raw = ("id,note,amount\r\n" + body).encode("utf-8")
    # The delimiter is given: it is sniffed from the first chunk, which in
    # production is never smaller than the sample the reader sniffs from.
    whole = preflight(raw, Format.CSV, delimiter=",")
    for size in (1, 2, 3, 7, 64):
        monkeypatch.setattr(safety, "_CHUNK", size)
        assert preflight(raw, Format.CSV, delimiter=",") == whole
    assert whole.rows == reader_rows(raw) == 300


def test_a_recorded_delimiter_is_used_and_an_absent_one_is_sniffed() -> None:
    raw = b"a;b;c\n1;2;3\n"
    assert preflight(raw, Format.CSV).columns == 3
    assert preflight(raw, Format.CSV, delimiter=",").columns == 1


def test_a_file_in_no_strict_encoding_still_passes_as_the_reader_reads_it() -> None:
    # 0x81 is undefined in cp1252 and invalid as UTF-8: the reader falls back
    # rather than refusing, and so does this.
    found = preflight(b"Name,Email\n\x81dam,adam@example.com\n", Format.CSV)
    assert found.safe and found.encoding == "latin-1"
    assert decode(b"Name,Email\n\x81dam,adam@example.com\n").encoding == "latin-1"


def test_an_empty_csv_is_not_a_safety_matter() -> None:
    # The reader refuses it with its own sentence; nothing here is unsafe.
    assert preflight(b"", Format.CSV).safe


# ── CSV: what is refused ─────────────────────────────────────────────────────


def test_a_csv_over_the_source_ceiling_is_refused() -> None:
    raw = b"Name\nAda\n"
    refusal = refused(raw, Format.CSV, limits(max_source_bytes=len(raw) - 1))
    assert refusal.code is RefusalCode.SOURCE_TOO_LARGE


def test_a_csv_over_the_decoded_budget_is_refused() -> None:
    raw = b"Name\n" + b"Ada Lovelace\n" * 5000
    refusal = refused(raw, Format.CSV, limits(max_decoded_string_bytes=len(raw)))
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE
    assert preflight(raw, Format.CSV, limits(max_decoded_string_bytes=len(raw) + 64)).safe


def test_one_astral_character_makes_a_whole_csv_four_bytes_wide() -> None:
    # The GR-352 shape exactly: a file that is ASCII but for one character.
    plain = b"Name\n" + b"Ada Lovelace\n" * 5000
    wide = plain + EMOJI.encode("utf-8") + b"\n"
    envelope = limits(max_decoded_string_bytes=2 * len(plain))
    assert preflight(plain, Format.CSV, envelope).widest_character == 1
    refusal = refused(wide, Format.CSV, envelope)
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE
    assert survey(wide, Format.CSV, envelope).widest_character == 4
    # What it is being protected from, stated as the arithmetic it is.
    assert sys.getsizeof(wide.decode("utf-8")) > 3.9 * len(plain)


def test_cp1252_bytes_that_look_like_an_emoji_are_not_costed_as_one() -> None:
    # F0 9F 98 80 is an emoji in UTF-8 and four ordinary characters in cp1252.
    # The 0xE9 on the last line settles it: the file is not UTF-8, the reader
    # will read it as cp1252, and nothing in it is astral.
    raw = b"Name\n" + b"\xf0\x9f\x98\x80" + b"a" * 4000 + b"\ncaf\xe9\n"
    assert decode(raw).encoding == "cp1252"
    envelope = limits(max_decoded_string_bytes=10_000)
    found = preflight(raw, Format.CSV, envelope)
    assert found.safe and found.encoding == "cp1252" and found.widest_character == 2


def test_csv_cells_and_columns_are_counted_and_refused() -> None:
    raw = ("a,b,c\n" + "1,2,3\n" * 9).encode()
    assert preflight(raw, Format.CSV, limits(max_cells=30)).cells == 30
    assert refused(raw, Format.CSV, limits(max_cells=29)).code is RefusalCode.TOO_MANY_CELLS
    wide = ("h\n" + ",".join("x" * 1 for _ in range(300)) + "\n").encode()
    refusal = refused(wide, Format.CSV, limits(max_columns=256))
    assert refusal.code is RefusalCode.TOO_MANY_COLUMNS
    assert refusal.observed == 300


def test_csv_rows_are_refused_only_when_the_caller_asks_for_a_ceiling() -> None:
    raw = b"Name\nAda\n\n  \nBo\nCy\n"
    assert preflight(raw, Format.CSV).rows == 3  # blank lines are not rows
    assert preflight(raw, Format.CSV, limits(max_rows=3)).safe
    assert refused(raw, Format.CSV, limits(max_rows=2)).code is RefusalCode.TOO_MANY_ROWS


def test_a_line_that_is_a_file_is_refused() -> None:
    from koras_import import MAX_CELL

    raw = b"Name\n" + b"a" * (MAX_CELL + 1) + b"\n"
    refusal = refused(raw, Format.CSV, PROVISIONAL_LIMITS)
    assert refusal.code is RefusalCode.LINE_TOO_LONG
    # And without a line ending at all, which is the shape that would
    # otherwise be carried from chunk to chunk for ever.
    assert (
        refused(b"a" * (MAX_CELL * 10), Format.CSV, PROVISIONAL_LIMITS).code
        is RefusalCode.LINE_TOO_LONG
    )


def test_an_unclosed_quotation_mark_is_a_refusal_and_not_a_csv_error() -> None:
    raw = b'Name,Note\nAda,"never closed\n' + b"more,and more\n" * 20_000
    refusal = refused(raw, Format.CSV, PROVISIONAL_LIMITS)
    assert refusal.code is RefusalCode.MALFORMED
    assert "unclosed quotation mark" in str(refusal)


# ── CSV: refused without the whole file ever being one string ────────────────


def test_the_csv_pass_never_holds_the_file_as_text() -> None:
    # Four megabytes of ASCII and one emoji. Decoded whole, as the reader
    # decodes it, that is a sixteen megabyte string. The pass is allowed all
    # of it here -- so this measures the scan, not the refusal -- and must
    # stay far below even the file's own size.
    raw = b"Name,Note\n" + b"Ada Lovelace,a note about her\n" * 140_000 + EMOJI.encode() + b"\n"
    assert len(raw) > 4 * 1024 * 1024
    envelope = limits(max_decoded_string_bytes=1 << 30, max_cells=10**9)
    tracemalloc.start()
    try:
        found = preflight(raw, Format.CSV, envelope)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert found.decoded_string_bytes > 16 * 1024 * 1024
    assert peak < 3 * 1024 * 1024


def test_a_refused_csv_is_refused_without_the_whole_file_decode(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from koras_import import reading

    def reached(_: bytes) -> None:
        raise AssertionError("the whole-file decode was reached by the safety pass")

    monkeypatch.setattr(reading, "decode", reached)
    raw = b"Name\n" + b"Ada Lovelace\n" * 5000 + EMOJI.encode() + b"\n"
    refusal = refused(raw, Format.CSV, limits(max_decoded_string_bytes=2 * len(raw)))
    assert refusal.code is RefusalCode.DECODED_TOO_LARGE


def test_a_format_nobody_reads_is_only_measured() -> None:
    assert preflight(b"[]", Format.JSON).safe
    assert refused(b"[]", Format.JSON, limits(max_source_bytes=1)).code is (
        RefusalCode.SOURCE_TOO_LARGE
    )
